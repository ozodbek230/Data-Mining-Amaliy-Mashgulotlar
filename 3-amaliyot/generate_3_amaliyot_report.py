# -*- coding: utf-8 -*-
"""
3-amaliy mashg'ulot bo'yicha to'liq Word hisoboti (3-amaliy_mashg'ulot_hisoboti.docx)
hamda Jupyter Notebook (3-amaliy_mashgulot.ipynb) yaratish skripti.
"""

import os
import json
import pandas as pd
import numpy as np
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

CURRENT_DIR = os.path.dirname(os.path.abspath(__file__))
CSV_PATH = os.path.join(CURRENT_DIR, "talabalar.csv")
df = pd.read_csv(CSV_PATH)

# ==============================================================================
# WORD HUJJATINI TAYYORLASH
# ==============================================================================
doc = Document()

for s in doc.sections:
    s.top_margin = Inches(0.79)
    s.bottom_margin = Inches(0.79)
    s.left_margin = Inches(0.85)
    s.right_margin = Inches(0.79)

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>'))

def set_table_borders(table, color="D3D3D3"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="4" w:space="0" w:color="{color}"/>
            <w:bottom w:val="single" w:sz="4" w:space="0" w:color="{color}"/>
            <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/>
            <w:insideV w:val="none"/>
            <w:left w:val="none"/>
            <w:right w:val="none"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

def format_row(row, is_header=False):
    for cell in row.cells:
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        for p in cell.paragraphs:
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.05
            for run in p.runs:
                run.font.name = "Times New Roman"
                run.font.size = Pt(9.5 if not is_header else 10)
                if is_header:
                    run.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)

def add_heading_1(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.bold = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(14)
    r.font.color.rgb = RGBColor(15, 23, 42)
    return p

def add_heading_2(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.bold = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor(30, 58, 138)
    return p

def add_body(text, bold_prefix=None, italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.bold = True
        r_pre.font.name = "Times New Roman"
        r_pre.font.size = Pt(11)
    
    r = p.add_run(text)
    r.italic = italic
    r.font.name = "Times New Roman"
    r.font.size = Pt(11)
    return p

# -----------------------------------------------------------------------------
# TITUL VARAQASI
# -----------------------------------------------------------------------------
p_vaz = doc.add_paragraph()
p_vaz.alignment = WD_ALIGN_PARAGRAPH.CENTER
r_v = p_vaz.add_run("O'ZBEKISTON RESPUBLIKASI\nOLIY TA'LIM, FAN VA INNOVATSIYALAR VAZIRLIGI\n")
r_v.bold = True
r_v.font.size = Pt(13)
r_v.font.name = "Times New Roman"

p_kaf = doc.add_paragraph()
p_kaf.alignment = WD_ALIGN_PARAGRAPH.CENTER
r_k = p_kaf.add_run("AXBOROT TEXNOLOGIYALARI VA SUN'IY INTELLEKT KAFEDRASI\n\"DATA MINING\" FANI")
r_k.font.size = Pt(12)
r_k.bold = True
r_k.font.name = "Times New Roman"

for _ in range(3):
    doc.add_paragraph()

p_title = doc.add_paragraph()
p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r_t = p_title.add_run("3-AMALIY MASHG'ULOT HISOBOTI\n")
r_t.bold = True
r_t.font.size = Pt(20)
r_t.font.name = "Times New Roman"
r_t.font.color.rgb = RGBColor(15, 23, 42)

p_sub = doc.add_paragraph()
p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r_sub = p_sub.add_run("Mavzu: Ma’lumotlar to‘plami, atributlar va ma’lumotlar turlarini aniqlash hamda tahlil qilish\n(«Talabalar» ma'lumotlar to'plami asosida barcha 4 ta topshiriq va nazorat savollari yechimi)")
r_sub.italic = True
r_sub.font.size = Pt(12.5)
r_sub.font.name = "Times New Roman"
r_sub.font.color.rgb = RGBColor(71, 85, 105)

for _ in range(4):
    doc.add_paragraph()

tbl_info = doc.add_table(rows=2, cols=2)
tbl_info.alignment = WD_TABLE_ALIGNMENT.CENTER
tbl_info.autofit = True

c00 = tbl_info.rows[0].cells[0]
c00.text = "Bajardi:"
c00.paragraphs[0].runs[0].bold = True
c00.paragraphs[0].runs[0].font.size = Pt(12)

c01 = tbl_info.rows[0].cells[1]
c01.text = "Talaba: ____________________\nGuruh: ____________________"
c01.paragraphs[0].runs[0].font.size = Pt(12)

c10 = tbl_info.rows[1].cells[0]
c10.text = "Qabul qildi:"
c10.paragraphs[0].runs[0].bold = True
c10.paragraphs[0].runs[0].font.size = Pt(12)

c11 = tbl_info.rows[1].cells[1]
c11.text = "O'qituvchi: ____________________\nBahosi: ____________________"
c11.paragraphs[0].runs[0].font.size = Pt(12)

for _ in range(4):
    doc.add_paragraph()

p_shahar = doc.add_paragraph()
p_shahar.alignment = WD_ALIGN_PARAGRAPH.CENTER
r_sh = p_shahar.add_run("Toshkent – 2026")
r_sh.font.size = Pt(12)
r_sh.font.name = "Times New Roman"

doc.add_page_break()

# -----------------------------------------------------------------------------
# ISHNING MAQSADI VA BERILGAN DATASET
# -----------------------------------------------------------------------------
add_heading_1("ISHNING MAQSADI VA TADQIQOT OB'EKTI")
add_body("Ishning maqsadi: Talabalarda ma’lumotlar to‘plami (dataset), obyekt va atribut tushunchalarini amalda ajrata olish, atributlarni turlarga (nominal, binar, ordinal, interval, nisbat) ajratish, ularning asosiy statistik tavsifini hisoblash hamda ma’lumotlar sifatidagi muammolarni (yetishmayotgan qiymatlar, outlierlar) aniqlash va bartaraf etish ko‘nikmasini shakllantirish.")

add_heading_2("Tadqiqotda ishlatiladigan «Talabalar» ma'lumotlar to'plami")
add_body("Hujjat shartida berilgan 10 ta talaba bo'yicha boshlang'ich jadval ko'rinishi:")

tbl_orig = doc.add_table(rows=1, cols=10)
tbl_orig.alignment = WD_TABLE_ALIGNMENT.CENTER
set_table_borders(tbl_orig)

orig_headers = ["ID", "Ism", "Jinsi", "Yoshi", "Kurs", "Yo‘nalish", "GPA", "Stipendiya", "Baho", "Tug‘.yil"]
for i, h in enumerate(orig_headers):
    tbl_orig.rows[0].cells[i].text = h
    set_cell_background(tbl_orig.rows[0].cells[i], "1E293B")
format_row(tbl_orig.rows[0], is_header=True)

for idx, (_, r) in enumerate(df.iterrows(), 1):
    row_c = tbl_orig.add_row().cells
    row_c[0].text = str(r["ID"])
    row_c[1].text = str(r["Ism"])
    row_c[2].text = str(r["Jinsi"])
    row_c[3].text = str(r["Yoshi"])
    row_c[4].text = str(r["Kurs"])
    row_c[5].text = str(r["Yo‘nalish"]) if pd.notna(r["Yo‘nalish"]) else "—"
    row_c[6].text = str(r["GPA"]) if pd.notna(r["GPA"]) else "—"
    row_c[7].text = str(r["Stipendiya"])
    row_c[8].text = str(r["Baho darajasi"])
    row_c[9].text = str(r["Tug‘ilgan yil"])
    bg = "F8FAFC" if idx % 2 == 1 else "FFFFFF"
    for cell in row_c:
        set_cell_background(cell, bg)
    format_row(tbl_orig.rows[-1])

doc.add_page_break()

# -----------------------------------------------------------------------------
# 1-TOPSHIRIQ
# -----------------------------------------------------------------------------
add_heading_1("1-TOPSHIRIQ. MA’LUMOTLAR TO‘PLAMINI TAVSIFLASH")

add_body("Jadval o'rganilib, berilgan savollarga quyidagi yozma javoblar tayyorlandi:")

add_body("Ushbu ma'lumotlar to'plamida 10 ta obyekt (satr) va 10 ta atribut (ustun) mavjud.", bold_prefix="1. Obyektlar va atributlar soni: ")

add_body("Ushbu to'plamda har bir obyekt (jadval qatori) alohida bir talabani, ya'ni uning shaxsiy, demografik hamda oliy ta'lim muassasasidagi akademik profilini ifodalaydi.", bold_prefix="2. Obyekt nimani ifodalaydi? ")

add_body("Bu vazifani 'ID' atributi bajaradi. Chunki ID faqat texnik maqsadlarda jadvaldagi yozuvlarni bir-biridan unikal ajratish (birlamchi kalit) uchun xizmat qiladi. U sun'iy belgilangan tartib raqami bo'lib, talabaning bilimi, yoshi yoki xususiyatlariga hech qanday bog'liqlikka ega emas, shuning uchun statistik tahlil yoki bashorat modellarida qatnashmaydi.", bold_prefix="3. Qaysi atribut faqat identifikatsiyalash uchun xizmat qiladi va tahlilda qatnashmaydi? Nima uchun? ")

add_body("Ushbu to'plam orqali quyidagi amaliy masalalarni tahlil qilish mumkin:\n"
         "  a) Talabaning o'rtacha o'zlashtirish bali (GPA) uning stipendiya olishiga qanchalik ta'sir qiladi (stipendiya oluvchi va olmaydigan talabalarning GPA farqi)?\n"
         "  b) Qaysi ta'lim yo'nalishi (Sun'iy intellekt, Dasturiy injiniring, Kompyuter injiniringi) talabalari eng yuqori o'rtacha GPA ko'rsatkichiga ega?",
         bold_prefix="4. Ma'lumotlar to'plami yordamida javob topish mumkin bo'lgan 2 ta amaliy savol: ")

# -----------------------------------------------------------------------------
# 2-TOPSHIRIQ
# -----------------------------------------------------------------------------
add_heading_1("2-TOPSHIRIQ. ATRIBUT TURLARINI ANIQLASH")

add_body("Har bir atributning o'lchov shkalasi, qiymatlar to'plami turi hamda tanlovning nazariy asoslanishi:")

tbl_attr = doc.add_table(rows=1, cols=4)
tbl_attr.alignment = WD_TABLE_ALIGNMENT.CENTER
set_table_borders(tbl_attr)

headers_attr = ["Atribut", "Turi (Shkala)", "Diskret / Uzluksiz", "Asoslash"]
for i, h in enumerate(headers_attr):
    tbl_attr.rows[0].cells[i].text = h
    set_cell_background(tbl_attr.rows[0].cells[i], "1E293B")
format_row(tbl_attr.rows[0], is_header=True)

attr_rows_data = [
    ("Jinsi", "Binar (simmetrik)", "Diskret", "Faqat 2 ta teng huquqli toifadan (Erkak / Ayol) iborat."),
    ("Yoshi", "Nisbat (Ratio)", "Diskret", "Haqiqiy mutlaq nol mavjud (tug'ilish payti). Karrali nisbat to'liq ma'noga ega (20 yosh 10 yoshdan 2 barobar katta)."),
    ("Kurs", "Ordinal (Tartib)", "Diskret", "Kurslar o'rtasida aniq ierarxiya va ketma-ketlik mavjud (1 < 2 < 3 < 4), lekin kurslar orasidagi bilim farqi son bilan ifodalanmaydi."),
    ("Yo‘nalish", "Nominal", "Diskret", "Mutaxassislik nomlari bo'lib, ularni matematik tartiblab yoki arifmetik amal bajarib bo'lmaydi."),
    ("GPA", "Nisbat (Ratio)", "Uzluksiz", "Mutlaq nol (0.0) mavjud, ma'lum oraliqda istalgan o'nli kasr qiymatni qabul qiladi, 4.0 ball 2.0 balldan 2 barobar yuqori."),
    ("Stipendiya", "Binar (asimmetrik)", "Diskret", "Faqat 2 ta holatdan ('Ha' yoki 'Yo‘q') iborat."),
    ("Baho darajasi", "Ordinal (Tartib)", "Diskret", "Darajalar o'rtasida tabiiy ierarxiya mavjud: Qoniqarli < Yaxshi < A’lo."),
    ("Tug‘ilgan yil", "Interval", "Diskret", "Taqvimiy shkala bo'lib, unda mutlaq nol yo'q (0-yil vaqtning boshlanishi emas). Yillar orasidagi farq ma'noli, ammo nisbati ma'nosiz.")
]

for idx, r in enumerate(attr_rows_data, 1):
    row_c = tbl_attr.add_row().cells
    row_c[0].text = r[0]
    row_c[1].text = r[1]
    row_c[2].text = r[2]
    row_c[3].text = r[3]
    bg = "F8FAFC" if idx % 2 == 1 else "FFFFFF"
    for cell in row_c:
        set_cell_background(cell, bg)
    format_row(tbl_attr.rows[-1])

add_heading_2("Qo'shimcha savollarga javoblar:")
add_body("Chunki talabalarning o'quv kursi shunchaki shartli belgi emas, balki aniq ta'lim bosqichlari va tajriba darajasini ifodalovchi ketma-ketlikka ega (1-kurs < 2-kurs < 3-kurs < 4-kurs). Unda kattalik-kichiklik munosabati mavjud bo'lgani uchun u ordinal shkalaga kiradi.", bold_prefix="• «Kurs» atributini nima uchun nominal emas, balki ordinal deb hisoblaymiz? ")

add_body("Ularning tub farqi mutlaq nol nuqtasi va karrali nisbatda namoyon bo'ladi:\n"
         "  - 'Tug‘ilgan yil' interval shkalada: Taqvimiy yillarda mutlaq fizik nol yo'q (milodiy 0-yil vaqt boshlanishi emas). Masalan, 2000-yilda tug'ilgan inson 1000-yilda tug'ilgandan 2 marta 'katta' deb bo'lmaydi.\n"
         "  - 'Yoshi' esa nisbat (ratio) shkalada: Unda haqiqiy mutlaq nol (insonning tug'ilish momenti = 0 yosh) mavjud. Shu sababli 20 yoshli talaba 10 yoshli boladan 2 barobar katta deb to'g'ri nisbat olish mumkin.",
         bold_prefix="• «Tug‘ilgan yil» va «Yoshi» atributlarining turlari nima bilan farq qiladi? ")

doc.add_page_break()

# -----------------------------------------------------------------------------
# 3-TOPSHIRIQ
# -----------------------------------------------------------------------------
add_heading_1("3-TOPSHIRIQ. ATRIBUTLARNING STATISTIK TAVSIFI")

add_body("Excel yoki Python yordamida hisoblangan sonli va toifaviy ko'rsatkichlar. Bo'sh qiymatlar hisobga olinmagan:")

add_heading_2("1. Sonli atributlar («Yoshi» va «GPA») statistikasi")
add_body("Izoh: 'Yoshi' atributida 9-satrda Otabekning yoshi 120 deb kiritilgan bo'lib, u g'ayrioddiy qiymat (outlier) hisoblanadi. Shuning uchun statistik ko'rsatkichlar ham asl xatolik bilan, ham xatolik olib tashlangan holatda hisoblandi:")

tbl_stats = doc.add_table(rows=1, cols=4)
tbl_stats.alignment = WD_TABLE_ALIGNMENT.CENTER
set_table_borders(tbl_stats)

h_stat = ["Ko'rsatkich", "Yoshi (Asl 120 bilan)", "Yoshi (Outliersiz - 9 ta)", "GPA (Mavjud 9 ta talaba)"]
for i, h in enumerate(h_stat):
    tbl_stats.rows[0].cells[i].text = h
    set_cell_background(tbl_stats.rows[0].cells[i], "1E293B")
format_row(tbl_stats.rows[0], is_header=True)

stats_table_data = [
    ("Minimum (Min)", "19 yosh", "19 yosh", "2.90 ball"),
    ("Maksimum (Max)", "120 yosh", "23 yosh", "4.50 ball"),
    ("Diapazon (Max - Min)", "101 yosh", "4 yosh", "1.60 ball"),
    ("O'rtacha qiymat (Mean)", "30.40 yosh", "20.44 yosh", "3.71 ball"),
    ("Mediana (Median)", "20.50 yosh", "20.00 yosh", "3.80 ball"),
    ("Moda (Mode)", "19 yosh (3 marta)", "19 yosh (3 marta)", "Mavjud emas (har biri 1 martadan)")
]

for idx, r in enumerate(stats_table_data, 1):
    row_c = tbl_stats.add_row().cells
    row_c[0].text = r[0]
    row_c[1].text = r[1]
    row_c[2].text = r[2]
    row_c[3].text = r[3]
    bg = "F8FAFC" if idx % 2 == 1 else "FFFFFF"
    for cell in row_c:
        set_cell_background(cell, bg)
    format_row(tbl_stats.rows[-1])

add_heading_2("2. Toifaviy atributlar («Yo‘nalish», «Stipendiya», «Baho darajasi») chastotasi va modasi")

tbl_cat = doc.add_table(rows=1, cols=4)
tbl_cat.alignment = WD_TABLE_ALIGNMENT.CENTER
set_table_borders(tbl_cat)

h_cat = ["Atribut nomi", "Qiymatlar va ularning uchrash soni (Chastota)", "Moda (Eng ko'p uchragan)", "Modaning soni"]
for i, h in enumerate(h_cat):
    tbl_cat.rows[0].cells[i].text = h
    set_cell_background(tbl_cat.rows[0].cells[i], "1E293B")
format_row(tbl_cat.rows[0], is_header=True)

cat_data = [
    ("Yo‘nalish", "Sun’iy intellekt: 4 ta (44.4%)\nDasturiy injiniring: 3 ta (33.3%)\nKompyuter injiniringi: 2 ta (22.2%)", "Sun’iy intellekt", "4 marta"),
    ("Stipendiya", "Ha: 5 ta (50.0%)\nYo‘q: 5 ta (50.0%)", "Bimodal (Ha va Yo‘q)", "teng 5 tadan"),
    ("Baho darajasi", "Yaxshi: 5 ta (50.0%)\nA’lo: 3 ta (30.0%)\nQoniqarli: 2 ta (20.0%)", "Yaxshi", "5 marta")
]

for idx, r in enumerate(cat_data, 1):
    row_c = tbl_cat.add_row().cells
    row_c[0].text = r[0]
    row_c[1].text = r[1]
    row_c[2].text = r[2]
    row_c[3].text = r[3]
    bg = "F8FAFC" if idx % 2 == 1 else "FFFFFF"
    for cell in row_c:
        set_cell_background(cell, bg)
    format_row(tbl_cat.rows[-1])

add_heading_2("3. «Baho darajasi» atributi uchun mediana va o'rtacha qiymat haqida")
add_body("Barcha 10 ta talabaning baho darajalarini o'sish tartibida joylashtiramiz:\n"
         "  [Qoniqarli, Qoniqarli, Yaxshi, Yaxshi, Yaxshi, Yaxshi, Yaxshi, A’lo, A’lo, A’lo]\n"
         "  Juft 10 ta element bo'lgani sababli, o'rtadagi 5- va 6-o'rindagi qiymatlar 'Yaxshi' hisoblanadi.\n"
         "  Demak, Mediana = 'Yaxshi'.",
         bold_prefix="• Mediana hisobi: ")

add_body("Chunki 'Baho darajasi' sonli emas, balki tartiblangan toifaviy (ordinal) o'zgaruvchidir. Matnli sifat toifalarini o'zaro qo'shish va ularni talabalar soniga bo'lish matematik va semantik jihatdan imkonsizdir (arifmetik amallarni faqat interval va nisbat shkalalarida bajarish mumkin).",
         bold_prefix="• Nima uchun o'rtacha qiymatni hisoblab bo'lmaydi? ")

doc.add_page_break()

# -----------------------------------------------------------------------------
# 4-TOPSHIRIQ
# -----------------------------------------------------------------------------
add_heading_1("4-TOPSHIRIQ. PYTHON YORDAMIDA TAHLIL VA MA’LUMOTLAR SIFATINI TEKSHIRISH")

add_body("Ma'lumotlar to'plami 'talabalar.csv' faylida saqlanib, Python (pandas) kutubxonasida tahlil qilindi:")

add_body("""import pandas as pd
df = pd.read_csv("talabalar.csv")
print(df.shape)            # (10, 10)
print(df.dtypes)           # ma'lumot turlari
print(df.describe())       # statistik ko'rsatkichlar
print(df.isnull().sum())   # bo'sh qiymatlar soni
print(df["Yo‘nalish"].value_counts())""", italic=True)

add_heading_2("Natijalar asosidagi tahlil va javoblar:")

add_body("Pandas aniqlagan turlar (int64, float64, object) bilan 2-topshiriqdagi nazariy turlar taqqoslanganda muhim farqlar mavjud:\n"
         "  • Pandas faqat xotirada saqlash turini biladi (masalan, ID, Kurs va Tug‘ilgan yilni bir xil int64 butun son deb ko'radi). Biroq nazariy jihatdan Kurs ordinal shkalada, Tug‘ilgan yil interval shkalada, ID esa shunchaki identifikator nomdir.\n"
         "  • Pandas Ism, Jinsi, Yo‘nalish, Stipendiya, Baho darajalarining barchasini umumiy 'object' (matn) deb qabul qiladi. Lekin ularning biri binar, biri nominal, biri ordinal shkalaga ega.",
         bold_prefix="a) Pandas turlari va nazariy turlar solishtirmasi: ")

add_body("Jadvalda 2 ta ustunda bittadan bo'sh qiymat (NaN) aniqlandi:\n"
         "  1. Bekzod (5-satr) GPA ko'rsatkichi bo'sh: Uni talabaning baho darajasiga mos ravishda to'ldirish maqsadga muvofiq. Bekzodning bahosi 'Yaxshi' bo'lgani sababli, 'Yaxshi' bahoga ega boshqa talabalarning o'rtacha GPA si (3.68) bilan yoki butun guruhning o'rtacha GPA si (3.71) bilan to'ldirish taklif etiladi.\n"
         "  2. Madina (10-satr) Yo'nalishi bo'sh: Bu toifaviy ustun bo'lgani uchun uni to'ldirishda eng ko'p uchragan moda qiymatidan ('Sun’iy intellekt') foydalaniladi.",
         bold_prefix="b) Bo'sh qiymatlar va ularni to'ldirish (Imputation) usullari: ")

add_body("9-satrdagi Otabekning yoshi 120 deb kiritilgan. Bu inson omili sababli ro'y bergan yaqqol xatolik (outlier) hisoblanadi (talabaning tug'ilgan yili 2005 bo'lib, uning yoshi aslida 20 yosh bo'lishi kerak edi):\n"
         "  • Ushbu 120 yoshli xato tufayli butun guruhning o'rtacha yoshi 30.4 yoshga chiqib ketgan.\n"
         "  • Agar ushbu g'ayrioddiy qiymat olib tashlansa (yoki to'g'ri 20 yoshga o'zgartirilsa), talabalarning haqiqiy o'rtacha yoshi 20.4 yoshni tashkil etadi.\n"
         "  • Ko'rinib turibdiki, bitta outlier o'rtacha qiymatni qariyb 10 yoshga (9.96 yoshga) sun'iy ravishda ko'tarib, statistika xolisligini butunlay buzadi.",
         bold_prefix="c) Shubhali (Outlier) qiymat va uning o'rtachaga ta'siri: ")

add_heading_2("Tozalangan va tiklangan yakuniy jadval (Cleaned Dataset)")
add_body("Yuqoridagi sifat muammolari (bo'sh qiymatlar to'ldirildi va 120 yosh to'g'rilandi) bartaraf etilgandan so'ng jadval ko'rinishi:")

tbl_clean = doc.add_table(rows=1, cols=10)
tbl_clean.alignment = WD_TABLE_ALIGNMENT.CENTER
set_table_borders(tbl_clean)

for i, h in enumerate(orig_headers):
    tbl_clean.rows[0].cells[i].text = h
    set_cell_background(tbl_clean.rows[0].cells[i], "1E293B")
format_row(tbl_clean.rows[0], is_header=True)

df_clean_rep = df.copy()
df_clean_rep.loc[df_clean_rep["ID"] == 5, "GPA"] = 3.68
df_clean_rep.loc[df_clean_rep["ID"] == 10, "Yo‘nalish"] = "Sun’iy intellekt"
df_clean_rep.loc[df_clean_rep["ID"] == 9, "Yoshi"] = 20

for idx, (_, r) in enumerate(df_clean_rep.iterrows(), 1):
    row_c = tbl_clean.add_row().cells
    row_c[0].text = str(r["ID"])
    row_c[1].text = str(r["Ism"])
    row_c[2].text = str(r["Jinsi"])
    row_c[3].text = str(r["Yoshi"])
    row_c[4].text = str(r["Kurs"])
    row_c[5].text = str(r["Yo‘nalish"])
    row_c[6].text = str(r["GPA"])
    row_c[7].text = str(r["Stipendiya"])
    row_c[8].text = str(r["Baho darajasi"])
    row_c[9].text = str(r["Tug‘ilgan yil"])
    bg = "F8FAFC" if idx % 2 == 1 else "FFFFFF"
    for cell in row_c:
        set_cell_background(cell, bg)
    format_row(tbl_clean.rows[-1])

doc.add_page_break()

# -----------------------------------------------------------------------------
# 5. NAZORAT SAVOLLARI VA XULOSA
# -----------------------------------------------------------------------------
add_heading_1("5. NAZORAT SAVOLLARI VA ULARNING JAVOBLARI")

qa_list = [
    ("1. Obyekt va atribut tushunchalarini misollar bilan tushuntiring.",
     "Obyekt — tahlil qilinayotgan mustaqil birlik, real voqelikdagi hodisa yoki shaxsdir (masalan: aniq bir talaba, mijoz, bank hisob raqami). "
     "Atribut — obyektning xossasi yoki belgisidir (masalan: talabaning yoshi, GPA ko'rsatkichi, tahsil olayotgan yo'nalishi)."),

    ("2. Interval va nisbat atributlarining asosiy farqi nimada?",
     "Asosiy farq mutlaq fizik nol (absolute zero) mavjudligida va karrali nisbat olinishidadir. "
     "Nisbat shkalasida mutlaq nol mavjud bo'lib, unda qiymatlarning o'zaro nisbati ma'noga ega (masalan, 20 yosh 10 yoshdan 2 marta katta). "
     "Interval shkalada esa mutlaq nol yo'q (masalan, 0°C harorat issiqlik yo'qligini anglatmaydi va 20°C 10°C dan 2 marta issiq deyilmaydi, faqat 10°C ga farq qiladi)."),

    ("3. Simmetrik va asimmetrik binar atributlarga misol keltiring.",
     "Simmetrik binar atributda ikkala holat ham bir xil ahamiyatga ega bo'ladi (masalan: Jinsi — Erkak / Ayol). "
     "Asimmetrik binar atributda esa bitta holat ikkinchisiga qaraganda ancha muhimroq hisoblanadi (masalan: Tibbiy tahlil natijasi — Kasallik bor / yo'q, yoki Kiberxavfsizlikda hujum bor / yo'q)."),

    ("4. Qaysi statistik ko'rsatkich g'ayrioddiy qiymatlarga (outlier) kamroq sezgir: o'rtacha yoki mediana? Nima uchun?",
     "Mediana outlierlarga nisbatan ancha chidamli (kamroq sezgir). Chunki arifmetik o'rtacha barcha sonlar yig'indisiga tayanadi va bitta keskin katta son (masalan, 120 yosh) o'rtachani birdaniga buzadi. "
     "Mediana esa faqatgina tartiblangan qatordagi o'rtada turgan elementning joylashuv pozitsiyasiga qaraydi va chetdagi anomal sonlarning hajmiga bog'liq emas."),

    ("5. Ma'lumotlar sifatini pasaytiruvchi asosiy omillarni sanab bering.",
     "1. Yetishmayotgan (bo'sh / Missing) ma'lumotlar;\n"
     "2. Shovqinli va anomaliya (Outliers) qiymatlar;\n"
     "3. Nomuvofiq va qarama-qarshi (Inconsistent) yozuvlar;\n"
     "4. Takroriy (Duplicate) ma'lumotlar;\n"
     "5. Inson omili tufayli yuzaga keladigan mexanik imlo va kiritish xatolari.")
]

for q, a in qa_list:
    add_body(a, bold_prefix=f"{q}\nJavob: ")

add_heading_1("6. XULOSA")
add_body("Atribut turini to'g'ri aniqlash ma'lumotlarni tahlil qilish usullari, vizualizatsiya shakli va mashinali o'qitish algoritmlarini xatosiz tanlashning eng muhim poydevori hisoblanadi. Agar toifaviy (ordinal yoki nominal) ma'lumotlarga arifmetik amallar qo'llanilsa yoki sonli ustunlardagi outlierlar tozalanmasa, olingan statistik xulosalar mutlaqo noto'g'ri bo'ladi. Ushbu amaliy mashg'ulot orqali ma'lumotlar sifatidagi nuqsonlarni (missing value, outlier) dastlabki bosqichdayoq aniqlash va ularni to'g'ri tiklash ko'nikmasi shakllantirildi.")

doc_path = os.path.join(CURRENT_DIR, "3-amaliy_mashg'ulot_hisoboti.docx")
doc.save(doc_path)
print(f"Word hisoboti saqlandi: {doc_path}")

# ==============================================================================
# JUPYTER NOTEBOOK (3-amaliy_mashgulot.ipynb)
# ==============================================================================
nb_cells = [
    {
        "cell_type": "markdown",
        "metadata": {},
        "source": [
            "# 3-Amaliy Mashg'ulot: Ma’lumotlar to‘plami, atributlar va ma’lumotlar turlarini aniqlash hamda tahlil qilish\n",
            "**Fan:** Data Mining  \n",
            "**Mavzu:** «Talabalar» ma'lumotlar to'plami asosida atribut turlarini ajratish, statistik ko'rsatkichlarni hisoblash va ma'lumotlar sifatini (outlier, missing values) tekshirish."
        ]
    },
    {
        "cell_type": "code",
        "execution_count": None,
        "metadata": {},
        "outputs": [],
        "source": [
            "# 1. Kutubxonalarni import qilish va datasetni yuklash\n",
            "import os\n",
            "import pandas as pd\n",
            "import numpy as np\n",
            "\n",
            "df = pd.read_csv('talabalar.csv')\n",
            "print('Dataset hajmi:', df.shape)\n",
            "display(df)"
        ]
    },
    {
        "cell_type": "markdown",
        "metadata": {},
        "source": [
            "## 1-topshiriq. Ma'lumotlar to'plamini tavsiflash\n",
            "1. **Obyektlar va atributlar:** 10 ta obyekt (satr) va 10 ta atribut (ustun) mavjud.\n",
            "2. **Obyekt:** Har bir satr alohida bir talabani ifodalaydi.\n",
            "3. **Identifikator:** `ID` atributi faqat tartib raqami bo'lib, tahlilda qatnashmaydi.\n",
            "4. **Amaliy savollar:**\n",
            "   - Stipendiya olish GPA ko'rsatkichiga bog'liqmi?\n",
            "   - Qaysi ta'lim yo'nalishi talabalarida eng yuqori GPA kuzatiladi?"
        ]
    },
    {
        "cell_type": "markdown",
        "metadata": {},
        "source": [
            "## 2-topshiriq. Atribut turlarini aniqlash jadvali\n",
            "| Atribut | Shkalasi | Turi | Asoslash |\n",
            "| :--- | :--- | :--- | :--- |\n",
            "| `Jinsi` | Binar (simmetrik) | Diskret | Faqat 2 ta toifa: Erkak / Ayol |\n",
            "| `Yoshi` | Nisbat (Ratio) | Diskret | Mutlaq nol bor, 20 yosh 10 yoshdan 2 marta katta |\n",
            "| `Kurs` | Ordinal | Diskret | 1 < 2 < 3 < 4 tartib va ierarxiya mavjud |\n",
            "| `Yo‘nalish` | Nominal | Diskret | Matematik tartiblanmaydigan soha nomlari |\n",
            "| `GPA` | Nisbat (Ratio) | Uzluksiz | Mutlaq nol bor (0.0), kasr sonli qiymatlar |\n",
            "| `Stipendiya` | Binar (asimmetrik) | Diskret | Faqat 2 ta holat: Ha / Yo'q |\n",
            "| `Baho darajasi`| Ordinal | Diskret | Qoniqarli < Yaxshi < A'lo darajasi |\n",
            "| `Tug‘ilgan yil`| Interval | Diskret | Taqvimiy shkala, mutlaq nol mavjud emas |"
        ]
    },
    {
        "cell_type": "markdown",
        "metadata": {},
        "source": [
            "## 3-topshiriq. Atributlarning statistik tavsifi"
        ]
    },
    {
        "cell_type": "code",
        "execution_count": None,
        "metadata": {},
        "outputs": [],
        "source": [
            "# Sonli atributlar statistikasi (Yoshi va GPA)\n",
            "yosh_all = df['Yoshi'].dropna()\n",
            "yosh_clean = df[df['Yoshi'] < 100]['Yoshi'].dropna()\n",
            "gpa = df['GPA'].dropna()\n",
            "\n",
            "stats_df = pd.DataFrame({\n",
            "    'Yoshi (120 bilan)': [yosh_all.min(), yosh_all.max(), yosh_all.max() - yosh_all.min(), round(yosh_all.mean(), 2), yosh_all.median(), yosh_all.mode()[0]],\n",
            "    'Yoshi (Tozalangan)': [yosh_clean.min(), yosh_clean.max(), yosh_clean.max() - yosh_clean.min(), round(yosh_clean.mean(), 2), yosh_clean.median(), yosh_clean.mode()[0]],\n",
            "    'GPA (9 ta talaba)': [gpa.min(), gpa.max(), round(gpa.max() - gpa.min(), 2), round(gpa.mean(), 2), gpa.median(), 'Moda yo\\'q']\n",
            "}, index=['Min', 'Max', 'Diapazon', 'O\\'rtacha', 'Mediana', 'Moda'])\n",
            "\n",
            "display(stats_df)"
        ]
    },
    {
        "cell_type": "code",
        "execution_count": None,
        "metadata": {},
        "outputs": [],
        "source": [
            "# Toifaviy atributlar: Yo'nalish, Stipendiya, Baho darajasi\n",
            "print('Yo\\'nalish chastotasi va modasi:')\n",
            "display(df['Yo‘nalish'].value_counts())\n",
            "print('Moda:', df['Yo‘nalish'].mode()[0])\n",
            "\n",
            "print('\\nStipendiya chastotasi:')\n",
            "display(df['Stipendiya'].value_counts())\n",
            "\n",
            "print('\\nBaho darajasi chastotasi va modasi:')\n",
            "display(df['Baho darajasi'].value_counts())\n",
            "print('Moda:', df['Baho darajasi'].mode()[0])\n",
            "print('Mediana: Yaxshi (chunki o\\'rtadagi 5 va 6-o\\'rindagi qiymat)')"
        ]
    },
    {
        "cell_type": "markdown",
        "metadata": {},
        "source": [
            "## 4-topshiriq. Python (pandas) orqali tahlil va sifatni tekshirish"
        ]
    },
    {
        "cell_type": "code",
        "execution_count": None,
        "metadata": {},
        "outputs": [],
        "source": [
            "# Pandas tekshiruvlari\n",
            "print('1. Hajmi:', df.shape)\n",
            "print('\\n2. Ma\\'lumot turlari:')\n",
            "display(df.dtypes)\n",
            "print('\\n3. Sonli ko\\'rsatkichlar:')\n",
            "display(df.describe())\n",
            "print('\\n4. Bo\\'sh qiymatlar soni:')\n",
            "display(df.isnull().sum())"
        ]
    },
    {
        "cell_type": "code",
        "execution_count": None,
        "metadata": {},
        "outputs": [],
        "source": [
            "# Bo'sh qiymatlarni to'ldirish va Outlierni to'g'irlash\n",
            "df_clean = df.copy()\n",
            "\n",
            "# 1. Bekzodning GPA sini 'Yaxshi' baholilar o'rtachasi bilan to'ldirish\n",
            "mean_yaxshi_gpa = df_clean[df_clean['Baho darajasi'] == 'Yaxshi']['GPA'].mean()\n",
            "df_clean['GPA'] = df_clean['GPA'].fillna(round(mean_yaxshi_gpa, 2))\n",
            "\n",
            "# 2. Madinaning Yo'nalishini moda bilan to'ldirish\n",
            "df_clean['Yo‘nalish'] = df_clean['Yo‘nalish'].fillna(df_clean['Yo‘nalish'].mode()[0])\n",
            "\n",
            "# 3. Otabekning yoshidagi 120 anomaliyasini 20 yoshga to'g'irlash (2005-yil tug'ilgan)\n",
            "df_clean.loc[df_clean['ID'] == 9, 'Yoshi'] = 20\n",
            "\n",
            "print('Tozalangan jadval:')\n",
            "display(df_clean)"
        ]
    }
]

nb_dict = {
    "cells": nb_cells,
    "metadata": {
        "kernelspec": {
            "display_name": "Python 3",
            "language": "python",
            "name": "python3"
        },
        "language_info": {
            "name": "python",
            "version": "3.12"
        }
    },
    "nbformat": 4,
    "nbformat_minor": 5
}

nb_path = os.path.join(CURRENT_DIR, "3-amaliy_mashgulot.ipynb")
with open(nb_path, "w", encoding="utf-8") as f:
    json.dump(nb_dict, f, ensure_ascii=False, indent=2)

print(f"Jupyter Notebook saqlandi: {nb_path}")
