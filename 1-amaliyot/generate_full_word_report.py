# -*- coding: utf-8 -*-
"""
Bitta to'liq Word faylga (1-amaliy_mashg'ulot_hisoboti.docx) 
barcha nazariya, kodlar, natijaviy jadvallar va xulosalarni joylash skripti.
"""

import os
import random
import numpy as np
import pandas as pd
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

np.random.seed(42)
random.seed(42)

CURRENT_DIR = os.path.dirname(os.path.abspath(__file__))
DATASET_DIR = os.path.join(CURRENT_DIR, "datasets")
os.makedirs(DATASET_DIR, exist_ok=True)

# 12 ta topshiriq ma'lumotlari va kodlari
TASKS_DATA = [
    {
        "id": 1,
        "soha": "Sog'liqni saqlash",
        "mavzu": "Shifoxonalar xarajatlari va ko'rsatilayotgan tibbiy xizmat sifati tahlili",
        "shart": "Ko'rsatilayotgan tibbiy xizmat sifatini aniqlash uchun shifoxonalarning sarflanadigan xarajat miqdorini: ovqatlanish, dorilar olish, tibbiy xizmat, boshqa xizmatlar haqidagi ma’lumotlar to’plangan jadvalni hosil qiling.",
        "fayl": "topshiriq_01_sogliqni_saqlash_xarajatlar.csv",
        "kod": """import pandas as pd
import numpy as np

np.random.seed(42)
n = 200

shifoxonalar = [f"Shifoxona_{i:03d}" for i in range(1, n + 1)]
hududlar = np.random.choice(["Toshkent sh.", "Samarqand", "Farg'ona", "Buxoro", "Andijon", "Qashqadaryo"], n)
ovqatlanish_xarajat = np.round(np.random.uniform(50, 300, n), 2)
dorilar_xarajat = np.round(np.random.uniform(120, 800, n), 2)
tibbiy_xizmat_xarajat = np.round(np.random.uniform(200, 1200, n), 2)
boshqa_xarajat = np.round(np.random.uniform(30, 250, n), 2)
jami_xarajat = np.round(ovqatlanish_xarajat + dorilar_xarajat + tibbiy_xizmat_xarajat + boshqa_xarajat, 2)
bemorlar_soni = np.random.randint(500, 5000, n)
sifat_bali = np.clip(np.round(55 + (jami_xarajat / bemorlar_soni) * 45 + np.random.normal(0, 4, n), 1), 50, 100)

df1 = pd.DataFrame({
    "Shifoxona_ID": shifoxonalar,
    "Hudud": hududlar,
    "Bemorlar_Soni": bemorlar_soni,
    "Ovqatlanish_mln": ovqatlanish_xarajat,
    "Dorilar_mln": dorilar_xarajat,
    "Tibbiy_Xizmat_mln": tibbiy_xizmat_xarajat,
    "Boshqa_Xarajat_mln": boshqa_xarajat,
    "Jami_Xarajat_mln": jami_xarajat,
    "Sifat_Bali_100": sifat_bali
})
df1.to_csv("datasets/topshiriq_01_sogliqni_saqlash_xarajatlar.csv", index=False)
print(df1.head())"""
    },
    {
        "id": 2,
        "soha": "Qurilish",
        "mavzu": "Qurilish loyihalari hajmi, pudratchi tajribasi, muddat va byudjet tahlili",
        "shart": "Loyiha hajmi va pudratchi tajribasi kabi omillarni hisobga olgan holda, qurilish loyihalarini belgilangan muddat va byudjet doirasida yakunlash ma’lumotlarini to’plang.",
        "fayl": "topshiriq_02_qurilish_loyihalari.csv",
        "kod": """import pandas as pd
import numpy as np

np.random.seed(42)
n = 200

loyihalar = [f"Loyiha_{i:03d}" for i in range(1, n + 1)]
bino_turi = np.random.choice(["Turar joy majmuasi", "Maktab / Bog'cha", "Tibbiyot markazi", "Biznes markazi"], n)
maydon_kv_m = np.random.randint(1500, 45000, n)
pudratchi_tajribasi_yil = np.random.randint(1, 28, n)
reja_byudjet_mlrd = np.round(maydon_kv_m * np.random.uniform(4.5, 6.8, n) / 1000, 2)
ortish_foizi = np.round(np.maximum(0, 22 - pudratchi_tajribasi_yil * 0.7 + np.random.normal(0, 4, n)), 1)
haqiqiy_xarajat_mlrd = np.round(reja_byudjet_mlrd * (1 + ortish_foizi / 100), 2)
reja_muddat_oy = np.random.randint(8, 36, n)
kechikish_oy = np.maximum(0, np.round((18 - pudratchi_tajribasi_yil * 0.6 + np.random.normal(0, 3, n)) / 4).astype(int))

df2 = pd.DataFrame({
    "Loyiha_ID": loyihalar,
    "Bino_Turi": bino_turi,
    "Maydon_kv_m": maydon_kv_m,
    "Pudratchi_Tajribasi_Yil": pudratchi_tajribasi_yil,
    "Reja_Byudjet_mlrd": reja_byudjet_mlrd,
    "Haqiqiy_Xarajat_mlrd": haqiqiy_xarajat_mlrd,
    "Reja_Muddat_Oy": reja_muddat_oy,
    "Kechikish_Oy": kechikish_oy,
    "Muvaffaqiyatli": np.where((kechikish_oy <= 1) & (ortish_foizi <= 5), "Ha", "Yo'q")
})
df2.to_csv("datasets/topshiriq_02_qurilish_loyihalari.csv", index=False)
print(df2.head())"""
    },
    {
        "id": 3,
        "soha": "Ta’lim",
        "mavzu": "Talabalar o'quv natijalari, davomat va kursda qolish xavfi prognozi",
        "shart": "O'quv natijalari, davomat va ijtimoiy-iqtisodiy omillarga asoslanib, kursdan-kursga qolish xavfi ostida bo’lgan talabalarning ma’lumotlarini shakllantiring.",
        "fayl": "topshiriq_03_talim_talabalar_xavfi.csv",
        "kod": """import pandas as pd
import numpy as np

np.random.seed(42)
n = 250

talabalar = [f"Talaba_{i:04d}" for i in range(1, n + 1)]
fakultet = np.random.choice(["KIF", "DIF", "AKT", "Iqtisodiyot"], n)
kurs = np.random.choice([1, 2, 3, 4], n)
davomat = np.round(np.clip(np.random.normal(82, 12, n), 40, 100), 1)
oraliq_baho = np.round(np.clip(np.random.normal(71, 14, n), 45, 100), 1)
mustaqil_ish_foiz = np.round(np.clip(davomat * 0.9 + np.random.normal(0, 7, n), 30, 100), 1)
ishlash_soati = np.random.choice([0, 10, 20, 30, 40], n, p=[0.4, 0.2, 0.2, 0.15, 0.05])

xavf_darajasi = []
for d, b, ish in zip(davomat, oraliq_baho, ishlash_soati):
    if d < 70 or b < 60 or (ish >= 30 and b < 70):
        xavf_darajasi.append("Yuqori xavf")
    elif d < 80 or b < 70:
        xavf_darajasi.append("O'rta xavf")
    else:
        xavf_darajasi.append("Past xavf")

df3 = pd.DataFrame({
    "Talaba_ID": talabalar,
    "Fakultet": fakultet,
    "Kurs": kurs,
    "Davomat_Foizi": davomat,
    "Oraliq_Baho": oraliq_baho,
    "Mustaqil_Ish_Foizi": mustaqil_ish_foiz,
    "Ishlash_Soati_Hafta": ishlash_soati,
    "Xavf_Darajasi": xavf_darajasi
})
df3.to_csv("datasets/topshiriq_03_talim_talabalar_xavfi.csv", index=False)
print(df3.head())"""
    },
    {
        "id": 4,
        "soha": "Sog’liqni saqlash",
        "mavzu": "Shifoxonadan keyin bemorlarning sog'lig'i va reabilitatsiyasi tahlili",
        "shart": "Shifoxonadan chiqqandan keyin bemorning sog'lig'i, davolanishga rioya qilish va keyingi parvarishlar tahlili haqidagi ma’lumotlarni to’plang.",
        "fayl": "topshiriq_04_bemor_reabilitatsiyasi.csv",
        "kod": """import pandas as pd
import numpy as np

np.random.seed(42)
n = 200

bemorlar = [f"Bemor_{i:04d}" for i in range(1, n + 1)]
yosh = np.random.randint(18, 85, n)
tashxis = np.random.choice(["Gipertoniya", "Qandli diabet", "Pnevmoniya", "Infarkt reabilitatsiya"], n)
dorilarga_rioya = np.round(np.clip(np.random.normal(80, 15, n), 30, 100), 1)
shifokor_nazorati = np.random.choice(["Ha", "Yo'q"], n, p=[0.75, 0.25])
parhez = np.random.choice(["A'lo", "Yaxshi", "Qoniqarsiz"], n, p=[0.35, 0.45, 0.20])
qayta_yotqizildi = ["Ha" if (d < 65 or k == "Yo'q") and random.random() < 0.6 else "Yo'q" 
                    for d, k in zip(dorilarga_rioya, shifokor_nazorati)]

df4 = pd.DataFrame({
    "Bemor_ID": bemorlar,
    "Yosh": yosh,
    "Asosiy_Tashxis": tashxis,
    "Dorilarga_Rioya_Foizi": dorilarga_rioya,
    "Shifokor_Nazoratiga_Keldi": shifokor_nazorati,
    "Parhez_Darajasi": parhez,
    "Qayta_Yotqizildi": qayta_yotqizildi
})
df4.to_csv("datasets/topshiriq_04_bemor_reabilitatsiyasi.csv", index=False)
print(df4.head())"""
    },
    {
        "id": 5,
        "soha": "Ta’lim",
        "mavzu": "O'qitish metodlarining samaradorligi va o'quvchilar faolligi tahlili",
        "shart": "Ta’lim texnikasi, o‘quvchilarning faolligi va ta’lim natijalari to‘g‘risidagi ma’lumotlardan foydalangan holda o‘quvchilar faoliyatini yaxshilashda turli o‘qitish usullarining samaradorligi ma’lumotlarini to’plang.",
        "fayl": "topshiriq_05_oqitish_usullari_samaradorligi.csv",
        "kod": """import pandas as pd
import numpy as np

np.random.seed(42)
n = 200

guruhlar = [f"Guruh_{i:03d}" for i in range(1, n + 1)]
fan = np.random.choice(["Matematika", "Dasturlash", "Fizika", "Sun'iy intellekt"], n)
metod = np.random.choice(["An'anaviy", "Interaktiv / Gamifikatsiya", "Loyiha asosida (PBL)", "Aralash (Blended)"], n)
faollik = [np.round(np.clip(np.random.normal(86 if "Interaktiv" in m or "PBL" in m else 67, 8), 40, 100), 1) for m in metod]
baho = [np.round(np.clip(np.random.normal(84 if "Interaktiv" in m or "PBL" in m else 71, 7), 50, 100), 1) for m in metod]

df5 = pd.DataFrame({
    "Guruh_ID": guruhlar,
    "Fan": fan,
    "Oqitish_Metodi": metod,
    "Darsdagi_Faollik_Foiz": faollik,
    "Yakuniy_Ortacha_Baho": baho,
    "Samaradorlik_Indeksi": np.round(np.array(baho) * 0.6 + np.array(faollik) * 0.4, 1)
})
df5.to_csv("datasets/topshiriq_05_oqitish_usullari_samaradorligi.csv", index=False)
print(df5.head())"""
    },
    {
        "id": 6,
        "soha": "Sog’liqni saqlash",
        "mavzu": "Epidemiologik ko'rsatkichlar va atrof-muhit ekologiyasi tahlili",
        "shart": "Epidemiologik ma'lumotlar, aholi salomatligi ko'rsatkichlari va atrof-muhit omillari asosida aholi salomatligi haqidagi ma’lumotlarni to’plang.",
        "fayl": "topshiriq_06_epidemiologiya_atrof_muhit.csv",
        "kod": """import pandas as pd
import numpy as np

np.random.seed(42)
n = 180

tumanlar = [f"Tuman_{i:03d}" for i in range(1, n + 1)]
aholi_ming = np.random.randint(40, 380, n)
havo_aqi = np.random.randint(35, 240, n)
yashil_maydon_foiz = np.round(np.clip(50 - (havo_aqi * 0.14) + np.random.normal(0, 4, n), 6, 60), 1)
sanoat_soni = np.random.randint(3, 55, n)
respirator_kasalliklar = np.round(14 + havo_aqi * 0.23 + sanoat_soni * 0.28 + np.random.normal(0, 3, n), 1)
umr_korsatkichi = np.round(np.clip(76.5 - (havo_aqi * 0.035), 65, 80), 1)

df6 = pd.DataFrame({
    "Tuman_ID": tumanlar,
    "Aholi_Ming": aholi_ming,
    "Havo_Sifati_AQI": havo_aqi,
    "Yashil_Maydon_Foiz": yashil_maydon_foiz,
    "Sanoat_Korxonalari": sanoat_soni,
    "Respirator_Kasalliklar_1000ga": respirator_kasalliklar,
    "Ortacha_Umr": umr_korsatkichi
})
df6.to_csv("datasets/topshiriq_06_epidemiologiya_atrof_muhit.csv", index=False)
print(df6.head())"""
    },
    {
        "id": 7,
        "soha": "Sport",
        "mavzu": "Futbol jamoalarining o'yin statistikasi, gollar va sovrinlar tahlili",
        "shart": "Futbol jamoasining g’alabalar soni, duranglar soni, mag’lubiyatlar soni, oxirgi besh yillikda erishgan sovrinlar soni, raqib darvozasiga kiritgan gollar soni, xavfli vaziyatlar soni, o’tkazib yuborgan gollari soni haqidagi ma’lumotlarni to’plang.",
        "fayl": "topshiriq_07_futbol_klublar_statistikasi.csv",
        "kod": """import pandas as pd
import numpy as np

np.random.seed(42)
n = 100

jamoalar = [f"Futbol_Klub_{i:03d}" for i in range(1, n + 1)]
galabalar = np.random.randint(4, 23, n)
duranglar = np.random.randint(3, 11, n)
maglubiyatlar = np.maximum(0, 30 - galabalar - duranglar)
ochkolar = galabalar * 3 + duranglar
urilgan_gollar = np.round(galabalar * 1.8 + duranglar * 0.8 + np.random.randint(0, 8, n)).astype(int)
otkazilgan_gollar = np.round(maglubiyatlar * 1.7 + duranglar * 0.7 + np.random.randint(0, 8, n)).astype(int)
xavfli_vaziyatlar = (urilgan_gollar * 3.4 + np.random.randint(20, 50, n)).astype(int)
sovrinlar = np.random.choice([0, 1, 2, 3, 5], n, p=[0.65, 0.18, 0.1, 0.05, 0.02])

df7 = pd.DataFrame({
    "Jamoa_Nomi": jamoalar,
    "Galabalar": galabalar,
    "Duranglar": duranglar,
    "Maglubiyatlar": maglubiyatlar,
    "Ochkolar": ochkolar,
    "Urilgan_Gollar": urilgan_gollar,
    "Otkazilgan_Gollar": otkazilgan_gollar,
    "Gollar_Farqi": urilgan_gollar - otkazilgan_gollar,
    "Xavfli_Vaziyatlar": xavfli_vaziyatlar,
    "Sovrinlar_Oxirgi_5_Yil": sovrinlar
})
df7.to_csv("datasets/topshiriq_07_futbol_klublar_statistikasi.csv", index=False)
print(df7.head())"""
    },
    {
        "id": 8,
        "soha": "Jamoat transporti",
        "mavzu": "Jamoat transporti yo'nalishlarining aniqligi, yuki va samaradorligi tahlili",
        "shart": "Umumiy tizim samaradorligini oshirish uchun jamoat transporti yo'nalishlarini aniqlik, yo'lovchilar yuki va infratuzilma sifati kabi ma’lumotlarni to’plang.",
        "fayl": "topshiriq_08_jamoat_transporti_samaradorligi.csv",
        "kod": """import pandas as pd
import numpy as np

np.random.seed(42)
n = 150

yonalishlar = [f"Yonalish_{i:03d}" for i in range(1, n + 1)]
transport = np.random.choice(["Avtobus", "Elektrobus", "Metro"], n, p=[0.55, 0.3, 0.15])
uzunlik_km = np.round(np.random.uniform(9, 42, n), 1)
kechikish_min = np.round(np.clip(np.random.exponential(3.2, n), 0.5, 22), 1)
yolovchilar_ming = np.round(uzunlik_km * np.random.uniform(0.5, 1.1, n), 1)
tigiz_yuklanish_foiz = np.round(np.clip(np.random.normal(83, 13, n), 45, 120), 1)
infratuzilma_10 = np.random.choice([5, 6, 7, 8, 9, 10], n, p=[0.1, 0.15, 0.3, 0.25, 0.15, 0.05])
reja_bajarilish = np.round(np.clip(100 - kechikish_min * 1.6, 60, 99.5), 1)

df8 = pd.DataFrame({
    "Yonalish_ID": yonalishlar,
    "Transport_Turi": transport,
    "Uzunlik_km": uzunlik_km,
    "Ortacha_Kechikish_min": kechikish_min,
    "Kunlik_Yolovchilar_Ming": yolovchilar_ming,
    "Tigiz_Yuklanish_Foiz": tigiz_yuklanish_foiz,
    "Infratuzilma_Sifati_10": infratuzilma_10,
    "Reja_Bajarilishi_Foiz": reja_bajarilish
})
df8.to_csv("datasets/topshiriq_08_jamoat_transporti_samaradorligi.csv", index=False)
print(df8.head())"""
    },
    {
        "id": 9,
        "soha": "Demografiya",
        "mavzu": "Demografik ko'rsatkichlar, ijtimoiy omillar va saylov faolligi tendensiyasi",
        "shart": "Demografik ma'lumotlar, ijtimoiy-iqtisodiy omillar va ovoz berishning oldingi tendensiyalari haqidagi ma’lumotlarni to’plang.",
        "fayl": "topshiriq_09_demografiya_saylov_tendensiyasi.csv",
        "kod": """import pandas as pd
import numpy as np

np.random.seed(42)
n = 150

okruglar = [f"Okrug_{i:03d}" for i in range(1, n + 1)]
aholi = np.random.randint(18000, 75000, n)
ortacha_yosh = np.round(np.random.uniform(27, 47, n), 1)
daromad_mln = np.round(np.random.uniform(2.8, 11.5, n), 2)
oliy_malumot_foiz = np.round(np.clip(np.random.normal(33, 9, n), 12, 70), 1)
ishsizlik_foiz = np.round(np.clip(np.random.normal(6.8, 2.2, n), 2.5, 15), 1)
davomat_foiz = np.round(np.clip(62 + oliy_malumot_foiz * 0.18 + np.random.normal(0, 3, n), 48, 95), 1)

df9 = pd.DataFrame({
    "Okrug_ID": okruglar,
    "Aholi_Soni": aholi,
    "Ortacha_Yosh": ortacha_yosh,
    "Ortacha_Daromad_mln": daromad_mln,
    "Oliy_Malumotlilar_Foiz": oliy_malumot_foiz,
    "Ishsizlik_Darajasi_Foiz": ishsizlik_foiz,
    "Saylov_Davomati_Foiz": davomat_foiz
})
df9.to_csv("datasets/topshiriq_09_demografiya_saylov_tendensiyasi.csv", index=False)
print(df9.head())"""
    },
    {
        "id": 10,
        "soha": "Sport",
        "mavzu": "Futbolchilar jarohatlari turlari, pozitsiyasi va maydon sharoitlari tahlili",
        "shart": "Futbol jamoasida o’yinchilarning jarohat turi, o'yinchi pozitsiyasi va atrof-muhit sharoitlariga qarab sport jarohatlari haqidagi ma’lumotlarni to’plang.",
        "fayl": "topshiriq_10_sportchilar_jaroxatlari.csv",
        "kod": """import pandas as pd
import numpy as np
import random

np.random.seed(42)
random.seed(42)
n = 220

oyinchilar = [f"Futbolchi_{i:04d}" for i in range(1, n + 1)]
yosh = np.random.randint(18, 35, n)
pozitsiya = np.random.choice(["Darvozabon", "Himoyachi", "Yarim himoyachi", "Hujumchi"], n, p=[0.1, 0.35, 0.35, 0.2])
jarohat = np.random.choice(["Mushak cho'zilishi", "Tizza bo'g'imi", "Boldir sinishi", "Pay cho'zilishi"], n)
maydon = np.random.choice(["Tabiiy maysa", "Sun'iy maysa", "Gibrid chim"], n)
sharoit = np.random.choice(["Quruq / quyoshli", "Yomg'irli / sirpanchiq", "Sovuq"], n)

hafta = []
for j in jarohat:
    if "Tizza" in j: hafta.append(random.randint(8, 34))
    elif "sinishi" in j: hafta.append(random.randint(10, 24))
    elif "Pay" in j: hafta.append(random.randint(3, 9))
    else: hafta.append(random.randint(1, 4))

df10 = pd.DataFrame({
    "Futbolchi_ID": oyinchilar,
    "Yosh": yosh,
    "Pozitsiya": pozitsiya,
    "Jarohat_Turi": jarohat,
    "Maydon_Qoplamasi": maydon,
    "Ob_Havo": sharoit,
    "Tiklanish_Muddati_Hafta": hafta,
    "Jarrohlik": ["Ha" if h > 12 else "Yo'q" for h in hafta]
})
df10.to_csv("datasets/topshiriq_10_sportchilar_jaroxatlari.csv", index=False)
print(df10.head())"""
    },
    {
        "id": 11,
        "soha": "Jamoat transporti",
        "mavzu": "Yo'lovchilarning tozalik, qulaylik va xizmat sifatidan qoniqish darajasi",
        "shart": "Tozalik, qulaylik va o'z vaqtida ishlash kabi omillarga asoslanib, jamoat transporti xizmatlaridan yo'lovchilarning qoniqishi ma’lumotlarini to’plang.",
        "fayl": "topshiriq_11_transport_yolovchilar_qoniqishi.csv",
        "kod": """import pandas as pd
import numpy as np

np.random.seed(42)
n = 250

sorov_id = [f"Respondent_{i:04d}" for i in range(1, n + 1)]
transport = np.random.choice(["Avtobus", "Metro", "Taksi"], n, p=[0.5, 0.35, 0.15])
tozalik = np.random.choice([1, 2, 3, 4, 5], n, p=[0.05, 0.15, 0.35, 0.30, 0.15])
qulaylik = np.random.choice([1, 2, 3, 4, 5], n, p=[0.1, 0.2, 0.3, 0.25, 0.15])
oz_vaqtida = np.random.choice([1, 2, 3, 4, 5], n, p=[0.08, 0.17, 0.30, 0.30, 0.15])
qoniqish_foiz = np.clip(np.round(((tozalik + qulaylik + oz_vaqtida) / 15) * 100 + np.random.normal(0, 3, n), 1), 20, 100)

df11 = pd.DataFrame({
    "Sorov_ID": sorov_id,
    "Transport_Turi": transport,
    "Tozalik_5": tozalik,
    "Qulaylik_5": qulaylik,
    "Oz_Vaqtida_5": oz_vaqtida,
    "Umumiy_Qoniqish_Foiz": qoniqish_foiz,
    "Tavsiya_Etadi": ["Ha" if q >= 70 else "Yo'q" for q in qoniqish_foiz]
})
df11.to_csv("datasets/topshiriq_11_transport_yolovchilar_qoniqishi.csv", index=False)
print(df11.head())"""
    },
    {
        "id": 12,
        "soha": "Demografiya",
        "mavzu": "Urbanizatsiya jarayonlari, aholi o'sishi va migratsiya tendensiyalari tahlili",
        "shart": "Aholi o'sishi, migratsiya shakllari va iqtisodiy ko'rsatkichlarni tahlil qilish orqali turli demografik guruhlarda urbanizatsiya tendensiyalari haqida ma’lumotlar to’plang.",
        "fayl": "topshiriq_12_urbanizatsiya_tendensiyalari.csv",
        "kod": """import pandas as pd
import numpy as np

np.random.seed(42)
n = 160

hududlar = [f"Hudud_{i:03d}" for i in range(1, n + 1)]
viloyat = np.random.choice(["Toshkent", "Samarqand", "Farg'ona", "Andijon", "Qashqadaryo", "Buxoro"], n)
aholi_ming = np.random.randint(50, 550, n)
osish_foiz = np.round(np.random.normal(1.8, 0.35, n), 2)
urbanizatsiya_foiz = np.round(np.clip(np.random.normal(48, 14, n), 18, 88), 1)
migratsiya_ming = np.round(np.random.normal(-0.4, 2.3, n), 2)
daromad_mln = np.round(np.clip(12 + (urbanizatsiya_foiz * 0.38) + np.random.normal(0, 3, n), 8, 55), 2)

df12 = pd.DataFrame({
    "Hudud_ID": hududlar,
    "Viloyat": viloyat,
    "Aholi_Ming": aholi_ming,
    "Yillik_Osish_Foiz": osish_foiz,
    "Urbanizatsiya_Darajasi_Foiz": urbanizatsiya_foiz,
    "Sof_Migratsiya_Ming": migratsiya_ming,
    "Jon_Boshiga_Daromad_mln": daromad_mln
})
df12.to_csv("datasets/topshiriq_12_urbanizatsiya_tendensiyalari.csv", index=False)
print(df12.head())"""
    }
]

# Word hujjatini yaratish
doc = Document()

# Sahifa o'lchamlari va chetki masofalari (Margins: 2 sm)
for s in doc.sections:
    s.top_margin = Inches(0.79)
    s.bottom_margin = Inches(0.79)
    s.left_margin = Inches(0.85)
    s.right_margin = Inches(0.79)

def add_title_page():
    # Vazirlik
    p_vazirlik = doc.add_paragraph()
    p_vazirlik.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_v = p_vazirlik.add_run("O'ZBEKISTON RESPUBLIKASI\nOLIY TA'LIM, FAN VA INNOVATSIYALAR VAZIRLIGI\n")
    r_v.bold = True
    r_v.font.size = Pt(13)
    r_v.font.name = "Times New Roman"

    p_kaf = doc.add_paragraph()
    p_kaf.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_k = p_kaf.add_run("AXBOROT-KOMMUNIKATSIYA TEXNOLOGIYALARI VA SUN'IY INTELLEKT KAFEDRASI\n\"DATA MINING\" FANI BO'YICHA")
    r_k.font.size = Pt(12)
    r_k.bold = True
    r_k.font.name = "Times New Roman"

    for _ in range(3):
        doc.add_paragraph()

    # Asosiy sarlavha
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_t = p_title.add_run("1-AMALIY MASHG'ULOT HISOBOTI\n")
    r_t.bold = True
    r_t.font.size = Pt(20)
    r_t.font.name = "Times New Roman"
    r_t.font.color.rgb = RGBColor(0, 51, 102)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_sub.add_run("Mavzu: Data Miningda ma'lumotlarni yig'ish, tozalash, Python (pandas, numpy) yordamida datasetlar yaratish va CSV formatida saqlash\n(Hujjatdagi barcha 12 ta topshiriqning to'liq yechimlari)")
    r_sub.italic = True
    r_sub.font.size = Pt(13)
    r_sub.font.name = "Times New Roman"
    r_sub.font.color.rgb = RGBColor(60, 60, 60)

    for _ in range(4):
        doc.add_paragraph()

    # Bajardi / Qabul qildi jadvali
    info_table = doc.add_table(rows=2, cols=2)
    info_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    info_table.autofit = True
    
    cell_00 = info_table.rows[0].cells[0]
    cell_00.text = "Bajardi:"
    cell_00.paragraphs[0].runs[0].bold = True
    cell_00.paragraphs[0].runs[0].font.size = Pt(12)
    
    cell_01 = info_table.rows[0].cells[1]
    cell_01.text = "Talaba: ____________________\nGuruh: ____________________"
    cell_01.paragraphs[0].runs[0].font.size = Pt(12)

    cell_10 = info_table.rows[1].cells[0]
    cell_10.text = "Qabul qildi:"
    cell_10.paragraphs[0].runs[0].bold = True
    cell_10.paragraphs[0].runs[0].font.size = Pt(12)

    cell_11 = info_table.rows[1].cells[1]
    cell_11.text = "O'qituvchi: ____________________\nBaholash:  ____________________"
    cell_11.paragraphs[0].runs[0].font.size = Pt(12)

    for _ in range(4):
        doc.add_paragraph()

    # Shahar va yil
    p_shahar = doc.add_paragraph()
    p_shahar.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sh = p_shahar.add_run("Toshkent – 2026")
    r_sh.font.size = Pt(12)
    r_sh.font.name = "Times New Roman"

    doc.add_page_break()

add_title_page()

# Mundarija / Reja
doc.add_heading("MUNDARIJA (REJA)", level=1)
reja_items = [
    "1. Ishning maqsadi va vazifalari",
    "2. Data Mining tushunchasi va ma'lumotlarni yig'ishning 11 ta bosqichi (Nazariy asoslar)",
    "3. Amaliy dasturlash vositalari (Python, Pandas, NumPy kutubxonalari)",
    "4. Amaliy topshiriqlarning to'liq dasturiy yechimlari (12 ta variant):",
    "   4.1. Sog'liqni saqlash: Shifoxona xarajatlari va tibbiy xizmat sifati",
    "   4.2. Qurilish: Loyiha hajmi, tajriba, muddat va byudjet",
    "   4.3. Ta'lim: Talabalar davomati, o'zlashtirishi va kursda qolish xavfi",
    "   4.4. Sog'liqni saqlash: Bemorlarning sog'lig'i va reabilitatsiya tahlili",
    "   4.5. Ta'lim: O'qitish metodlarining samaradorligi va darsdagi faollik",
    "   4.6. Sog'liqni saqlash: Epidemiologik ko'rsatkichlar va atrof-muhit ekologiyasi",
    "   4.7. Sport: Futbol jamoalari va klublarining to'liq statistikasi",
    "   4.8. Jamoat transporti: Yo'nalishlar samaradorligi va yo'lovchilar yuki",
    "   4.9. Demografiya: Ijtimoiy-iqtisodiy omillar va saylov faolligi tendensiyasi",
    "   4.10. Sport: Futbolchilar jarohatlari tahlili va maydon sharoitlari",
    "   4.11. Jamoat transporti: Yo'lovchilarning tozalik va xizmatdan qoniqishi",
    "   4.12. Demografiya: Urbanizatsiya tendensiyalari va migratsiya jarayonlari",
    "5. Amaliy mashg'ulot bo'yicha umumiy xulosa",
    "6. Foydalanilgan adabiyotlar ro'yxati"
]
for item in reja_items:
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(item)
    run.font.name = "Times New Roman"
    run.font.size = Pt(11)
    if not item.startswith("   "):
        run.bold = True

doc.add_page_break()

# 1. Ishning maqsadi
doc.add_heading("1. Ishning maqsadi va vazifalari", level=1)
p_m = doc.add_paragraph(
    "Data Mining (Ma'lumotlarni intellektual tahlili) sohasida ma'lumotlarni yig'ish, generatsiya qilish va tayyorlash bosqichlarini amaliy o'rganish. "
    "Python dasturlash tilining pandas va numpy kutubxonalaridan foydalangan holda, real hayotiy va iqtisodiy jarayonlarni (tibbiyot, qurilish, ta'lim, sport, "
    "jamoat transporti va demografiya) ifodalovchi ma'lumotlar to'plamlarini (DataFrames) avtomatlashtirilgan tarzda yaratish, ularga dastlabki ishlov berish "
    "hamda standart .csv fayllariga eksport qilish ko'nikmalarini shakllantirish."
)
p_m.paragraph_format.line_spacing = 1.15

# 2. Nazariy qism
doc.add_heading("2. Data Mining tushunchasi va ma'lumotlarni yig'ishning 11 ta bosqichi", level=1)
p_naz = doc.add_paragraph(
    "Data Mining — bu katta hajmdagi xom ma'lumotlar bazalarini tahlil qilish, ulardan yashirin qonuniyatlar, foydali bilimlar va bashoratli ma'lumotlarni "
    "ajratib olish jarayonidir. Xom ma'lumotlarni to'g'ridan-to'g'ri tahlil qilib bo'lmaydi, chunki ular to'liq emas, shovqinli (xatoliklarga ega) va noaniq bo'ladi. "
    "Shu sababli ma'lumotlarni tayyorlash quyidagi 11 ta muhim bosqichdan iborat bo'ladi:"
)
p_naz.paragraph_format.line_spacing = 1.15

bosqichlar = [
    ("1. Muammo ta’rifi", "Loyiha maqsadlarini aniqlash, tahlil orqali qaysi muammo yoki savolga javob izlanayotganini aniq belgilab olish."),
    ("2. Ma’lumotlarni oʻrganish", "Soha (domen) xususiyatlarini tushunish, potentsial manbalarni o'rganish va kerakli parametrlar turlarini aniqlash."),
    ("3. Ma’lumotlarni tanlash", "Mavjud omborlardan tahlil maqsadiga mos, yuqori sifatli va to'liq ma'lumot manbalarini ajratib olish."),
    ("4. Ma’lumotlarni tozalash", "Yetishmayotgan (NaN / None) qiymatlarni to'ldirish yoki o'chirish, xato va nomuvofiq qiymatlarni bartaraf etish."),
    ("5. Ma’lumotlarni oldindan ishlov berish", "Raqamli ustunlarni normallashtirish (Min-Max yoki Standartlashtirish), matnli ustunlarni One-Hot yoki Label encoding usulida kodlash."),
    ("6. Ma’lumotlar integratsiyasi", "Turli jadvallar va manbalardan olingan ma'lumotlarni yagona yaxlit ma'lumotlar to'plamiga birlashtirish."),
    ("7. Ma’lumotlarni qisqartirish", "Hajmi juda katta bo'lgan ma'lumotlarning o'lchamini kamaytirish (Feature Selection), ahamiyatsiz ustunlarni chiqarib tashlash."),
    ("8. Tanlanma olish (Sampling)", "Katta ma'lumotlar bazasidan butun to'plam xususiyatlarini o'zida aks ettiruvchi vakolatli kichik tanlanma ajratib olish."),
    ("9. Ma’lumotlarni oʻzgartirish (Transformation)", "Algoritmlarga moslashtirish maqsadida qo'shimcha yangi o'zgaruvchilarni hisoblab chiqarish va agregatsiyalash."),
    ("10. Ma’lumotlarni boʻlish", "Model sifatini baholash uchun datasetni o'qitish (train, odatda 70-80%) va test (20-30%) qismlariga ajratish."),
    ("11. Ma’lumotlarni saqlash", "Tayyorlangan toza ma'lumotlarni keyingi tahlillar va mashinali o'rganish modellariga uzatish uchun CSV, Parquet yoki ma'lumotlar bazasida saqlash.")
]

for b_nomi, b_matn in bosqichlar:
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    run_b = p.add_run(f"• {b_nomi}: ")
    run_b.bold = True
    run_b.font.color.rgb = RGBColor(0, 51, 102)
    p.add_run(b_matn)

# 3. Dasturiy ta'minot
doc.add_heading("3. Amaliy dasturlash vositalari: Pandas va NumPy kutubxonalari", level=1)
p_kutub = doc.add_paragraph(
    "Amaliy mashg'ulotni bajarishda Python tilining eng yetakchi ilmiy kutubxonalari ishlatildi:\n"
    "• NumPy — ko'p o'lchovli massivlar va tasodifiy sonlar generatorlari (np.random.normal, uniform, randint) bilan ishlash;\n"
    "• Pandas — Data Miningning asosiy vositasi bo'lib, ikki o'lchovli jadvallar (DataFrame) yaratish, ustunlar bo'yicha matematik hisob-kitoblar o'tkazish hamda .to_csv() va pd.read_csv() funksiyalari orqali fayllar bilan ishlashni ta'minlaydi."
)
p_kutub.paragraph_format.line_spacing = 1.15

doc.add_page_break()

# 4. Topshiriqlar yechimi
doc.add_heading("4. Amaliy topshiriqlarning to'liq dasturiy yechimlari", level=1)

for t in TASKS_DATA:
    t_id = t["id"]
    t_soha = t["soha"]
    t_mavzu = t["mavzu"]
    t_shart = t["shart"]
    t_fayl = t["fayl"]
    t_kod = t["kod"]
    
    # Faylni o'qish (allaqachon yaratilgan CSV dan)
    csv_path = os.path.join(DATASET_DIR, t_fayl)
    df = pd.read_csv(csv_path)

    doc.add_heading(f"4.{t_id}. {t_soha}: {t_mavzu}", level=2)
    
    # Shart
    p_sh = doc.add_paragraph()
    p_sh.paragraph_format.line_spacing = 1.15
    r_sh_lbl = p_sh.add_run("Topshiriq sharti: ")
    r_sh_lbl.bold = True
    p_sh.add_run(t_shart)

    p_f = doc.add_paragraph()
    p_f.add_run("Hosil qilingan fayl: ").bold = True
    p_f.add_run(f"datasets/{t_fayl} (Jami {len(df)} ta yozuv, {len(df.columns)} ta parametr/ustun)")

    # Python dasturiy kodi
    doc.add_paragraph("Python dasturiy kodi:").bold = True
    
    # Kod bloki uchun jadval (ramkali va och kulrang fonli)
    code_table = doc.add_table(rows=1, cols=1)
    code_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = code_table.rows[0].cells[0]
    cell.text = t_kod
    # Shading kulrang fon
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F4F6F9"/>')
    cell._tc.get_or_add_tcPr().append(shd)
    for p_c in cell.paragraphs:
        p_c.paragraph_format.space_before = Pt(3)
        p_c.paragraph_format.space_after = Pt(3)
        for r_c in p_c.runs:
            r_c.font.name = "Consolas"
            r_c.font.size = Pt(8.5)
            r_c.font.color.rgb = RGBColor(30, 30, 30)

    doc.add_paragraph()

    # Natijaviy jadval namunasi
    p_tbl_lbl = doc.add_paragraph("Hosil qilingan ma'lumotlar to'plamining dastlabki 5 ta qatori (head):")
    p_tbl_lbl.italic = True
    
    sample_df = df.head(5)
    cols = list(sample_df.columns)
    table = doc.add_table(rows=len(sample_df) + 1, cols=len(cols))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True

    # Sarlavha qatori
    hdr_cells = table.rows[0].cells
    for c_i, c_nomi in enumerate(cols):
        hdr_cells[c_i].text = str(c_nomi)
        shd_hdr = parse_xml(f'<w:shd {nsdecls("w")} w:fill="1B365D"/>')
        hdr_cells[c_i]._tc.get_or_add_tcPr().append(shd_hdr)
        for p in hdr_cells[c_i].paragraphs:
            for r in p.runs:
                r.bold = True
                r.font.name = "Times New Roman"
                r.font.size = Pt(8)
                r.font.color.rgb = RGBColor(255, 255, 255)

    # Qatorlar
    for r_i, (_, row) in enumerate(sample_df.iterrows()):
        row_cells = table.rows[r_i + 1].cells
        # Har ikkinchi qatorga yengil fon
        row_fill = "F9FAFC" if r_i % 2 == 1 else "FFFFFF"
        for c_i, c_nomi in enumerate(cols):
            row_cells[c_i].text = str(row[c_nomi])
            shd_row = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{row_fill}"/>')
            row_cells[c_i]._tc.get_or_add_tcPr().append(shd_row)
            for p in row_cells[c_i].paragraphs:
                for r in p.runs:
                    r.font.name = "Times New Roman"
                    r.font.size = Pt(8)

    # Statistik qisqacha sharh
    p_stat = doc.add_paragraph()
    p_stat.paragraph_format.space_before = Pt(6)
    p_stat.add_run("Tahlil natijasi: ").bold = True
    p_stat.add_run(
        f"Ushbu model doirasida {len(df)} ta ob'ektning {len(df.columns)} xil parametri bo'yicha matematik qonuniyatlar aks ettirildi. "
        f"Kategorik va sonli o'zgaruvchilar to'liq shakllantirilib, ma'lumotlar CSV formatida saqlandi."
    )

    doc.add_paragraph() # Bo'sh joy

doc.add_page_break()

# 5. Xulosa
doc.add_heading("5. Amaliy mashg'ulot bo'yicha umumiy xulosa", level=1)
p_xul = doc.add_paragraph(
    "Mazkur 1-amaliy mashg'ulot davomida Data Mining fanining eng asosiy poydevori hisoblangan "
    "ma'lumotlarni yig'ish, tozalash, transformatsiya qilish va ularga dastlabki ishlov berish bosqichlari to'liq o'rganildi.\n\n"
    "Bajarilgan ishlar yakuni bo'yicha quyidagi xulosalarga kelindi:\n"
    "1. Zamonaviy Data Mining loyihalarida vaqt va resurslarning 70-80% qismi ma'lumotlarni to'g'ri yig'ish va tozalash bosqichlariga sarflanadi. Ma'lumot sifatsiz bo'lsa, har qanday murakkab model ham noto'g'ri natija beradi ('Garbage In, Garbage Out' tamoyili).\n"
    "2. Python dasturlash tilining pandas va numpy kutubxonalari real jarayonlar modellarini tuzish va yuqori o'lchovli ma'lumotlar matritsalarini hosil qilish uchun qulay va yuqori samarali vosita ekanligi isbotlandi.\n"
    "3. Amaliy topshiriqdagi 12 ta turli soha (sog'liqni saqlash, qurilish, ta'lim, sport, jamoat transporti va demografiya) uchun ko'p parametrli sintetik datasetlar yaratildi. Har bir dataset mustaqil CSV fayllariga eksport qilindi va keyingi amaliy mashg'ulotlarda klasterlash, tasniflash (klassifikatsiya) hamda regressiya algoritmlarida foydalanishga to'liq tayyor holatga keltirildi."
)
p_xul.paragraph_format.line_spacing = 1.15

# 6. Foydalanilgan adabiyotlar
doc.add_heading("6. Foydalanilgan adabiyotlar ro'yxati", level=1)
adabiyotlar = [
    "1. Han, J., Kamber, M., & Pei, J. (2011). Data Mining: Concepts and Techniques. Morgan Kaufmann Publishers.",
    "2. McKinney, W. (2022). Python for Data Analysis: Data Wrangling with Pandas, NumPy, and Jupyter. O'Reilly Media.",
    "3. Tan, P.-N., Steinbach, M., Karpatne, A., & Kumar, V. (2018). Introduction to Data Mining. Pearson.",
    "4. Raschka, S., & Mirjalili, V. (2019). Python Machine Learning. Packt Publishing.",
    "5. Pandas documentation: https://pandas.pydata.org/docs/",
    "6. NumPy documentation: https://numpy.org/doc/"
]
for a in adabiyotlar:
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(a)
    r.font.name = "Times New Roman"
    r.font.size = Pt(11)

out_word_path = os.path.join(CURRENT_DIR, "1-amaliy_mashg'ulot_hisoboti.docx")
doc.save(out_word_path)
print(f"[SUCCESS] Yangilangan to'liq hisobot saqlandi: {out_word_path} (Hajmi: {os.path.getsize(out_word_path)} bayt)")
