# -*- coding: utf-8 -*-
"""
2-amaliy mashg'ulot bo'yicha to'liq Word hisoboti (2-amaliy_mashg'ulot_hisoboti.docx)
hamda Jupyter Notebook (2-amaliy_mashgulot.ipynb) yaratish skripti.
"""

import os
import json
import random
import numpy as np
import pandas as pd
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

np.random.seed(42)
random.seed(42)

CURRENT_DIR = os.path.dirname(os.path.abspath(__file__))
DATASETS_DIR = os.path.join(CURRENT_DIR, "datasets")

# Datasetlarni yuklash
df_football = pd.read_csv(os.path.join(DATASETS_DIR, "topshiriq_07_futbol_klublar_statistikasi.csv"))
df_talaba = pd.read_csv(os.path.join(DATASETS_DIR, "topshiriq_03_talim_talabalar_xavfi.csv"))

# Word Hujjatini yaratish
doc = Document()

for s in doc.sections:
    s.top_margin = Inches(0.79)
    s.bottom_margin = Inches(0.79)
    s.left_margin = Inches(0.85)
    s.right_margin = Inches(0.79)

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>'))

def set_table_borders(table, color="D3D3D3"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="4" w:space="0" w:color="{color}"/>
            <w:bottom w:val="single" w:sz="4" w:space="0" w:color="{color}"/>
            <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/>
            <w:insideV w:val="none"/>
            <w:left w:val="none"/>
            <w:right w:val="none"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

def format_row(row, is_header=False):
    for cell in row.cells:
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        for p in cell.paragraphs:
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.05
            for run in p.runs:
                run.font.name = "Times New Roman"
                run.font.size = Pt(9.5 if not is_header else 10)
                if is_header:
                    run.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)

def add_heading_1(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.bold = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(14)
    r.font.color.rgb = RGBColor(15, 23, 42)
    return p

def add_heading_2(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.bold = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor(30, 58, 138)
    return p

def add_body(text, bold_prefix=None, italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.bold = True
        r_pre.font.name = "Times New Roman"
        r_pre.font.size = Pt(11)
    
    r = p.add_run(text)
    r.italic = italic
    r.font.name = "Times New Roman"
    r.font.size = Pt(11)
    return p

# ==============================================================================
# TITUL VARAQASI
# ==============================================================================
p_vaz = doc.add_paragraph()
p_vaz.alignment = WD_ALIGN_PARAGRAPH.CENTER
r_v = p_vaz.add_run("O'ZBEKISTON RESPUBLIKASI\nOLIY TA'LIM, FAN VA INNOVATSIYALAR VAZIRLIGI\n")
r_v.bold = True
r_v.font.size = Pt(13)
r_v.font.name = "Times New Roman"

p_kaf = doc.add_paragraph()
p_kaf.alignment = WD_ALIGN_PARAGRAPH.CENTER
r_k = p_kaf.add_run("AXBOROT TEXNOLOGIYALARI VA SUN'IY INTELLEKT KAFEDRASI\n\"DATA MINING\" FANI")
r_k.font.size = Pt(12)
r_k.bold = True
r_k.font.name = "Times New Roman"

for _ in range(3):
    doc.add_paragraph()

p_title = doc.add_paragraph()
p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r_t = p_title.add_run("2-AMALIY MASHG'ULOT HISOBOTI\n")
r_t.bold = True
r_t.font.size = Pt(20)
r_t.font.name = "Times New Roman"
r_t.font.color.rgb = RGBColor(15, 23, 42)

p_sub = doc.add_paragraph()
p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r_sub = p_sub.add_run("Mavzu: Ma'lumotlar to'plami va uning atributlari. Metama'lumotlar\n(Tanlangan 2 ta CSV dataset bo'yicha barcha 7 ta amaliy vazifaning to'liq tahlili)")
r_sub.italic = True
r_sub.font.size = Pt(12.5)
r_sub.font.name = "Times New Roman"
r_sub.font.color.rgb = RGBColor(71, 85, 105)

for _ in range(4):
    doc.add_paragraph()

# Bajardi / Qabul qildi jadvali
tbl_info = doc.add_table(rows=2, cols=2)
tbl_info.alignment = WD_TABLE_ALIGNMENT.CENTER
tbl_info.autofit = True

c00 = tbl_info.rows[0].cells[0]
c00.text = "Bajardi:"
c00.paragraphs[0].runs[0].bold = True
c00.paragraphs[0].runs[0].font.size = Pt(12)

c01 = tbl_info.rows[0].cells[1]
c01.text = "Talaba: ____________________\nGuruh: ____________________"
c01.paragraphs[0].runs[0].font.size = Pt(12)

c10 = tbl_info.rows[1].cells[0]
c10.text = "Qabul qildi:"
c10.paragraphs[0].runs[0].bold = True
c10.paragraphs[0].runs[0].font.size = Pt(12)

c11 = tbl_info.rows[1].cells[1]
c11.text = "O'qituvchi: ____________________\nBahosi: ____________________"
c11.paragraphs[0].runs[0].font.size = Pt(12)

for _ in range(4):
    doc.add_paragraph()

p_shahar = doc.add_paragraph()
p_shahar.alignment = WD_ALIGN_PARAGRAPH.CENTER
r_sh = p_shahar.add_run("Toshkent – 2026")
r_sh.font.size = Pt(12)
r_sh.font.name = "Times New Roman"

doc.add_page_break()

# ==============================================================================
# NAZARIY QISM
# ==============================================================================
add_heading_1("I. NAZARIY QISM: ASOSIY TUSHUNCHALAR VA TA'RIFLAR")

add_body("Data Mining va ma'lumotlar tahlilida har qanday tadqiqot ma'lumotlar to'plamining strukturasini, undagi obyektlar va atributlarning tabiatini to'liq anglashdan boshlanadi. 2-amaliy mashg'ulot quyidagi asosiy tushunchalarni o'rganishga qaratilgan:")

concepts = [
    ("Obyekt (Record, Instance, Tuple): ", "Ma'lumotlar bazasi yoki jadvalidagi bitta yaxlit birlik, yozuv, holat, misol yoki jadval qatoridir. Har bir obyekt ma'lum bir real olam elementi (masalan, aniq bir futbol klubi yoki aniq bir talaba) haqidagi axborotni o'zida ifodalaydi."),
    ("Atribut (Feature, Dimension, Column): ", "Obyektni tavsiflovchi xususiyat, xossa yoki belgi. Jadval ko'rinishida atributlar ustunlar sifatida namoyon bo'ladi."),
    ("O'zgaruvchi (Variable): ", "O'rganilayotgan barcha obyektlar uchun umumiy xususiyat yoki xossa bo'lib, uning qiymati obyektdan obyektga farq qilishi mumkin."),
    ("Umumiy populyatsiya (Bosh to'plam, Population): ", "Tadqiqotchi yoki modelni qiziqtiradigan, o'rganilayotgan obyektlarning barcha nazariy va amaliy butun majmuasidir."),
    ("Tanlanma (Sample): ", "Umumiy populyatsiyaning xossalari va xususiyatlari to'g'risida tadqiqot olib borish va umumiy xulosalar chiqarish uchun ma'lum bir tartibda ajratib olingan vakillik (reprezentativ) qismi."),
    ("Statistika (Statistics): ", "Namunaning (tanlanmaning) raqamli xarakteristikalari, o'rtacha qiymat, dispersiya, mediana, korrelyatsiya ko'rsatkichlari."),
    ("Gipoteza (Hypothesis): ", "O'rganilayotgan obyektlar to'plamining parametrlari, atributlari orasidagi bog'liqliklar haqidagi ilmiy yoki amaliy taxmin bo'lib, u haqiqiy ma'lumotlar asosida tekshiriladi."),
    ("Diskret ma'lumotlar: ", "Atribut qiymatlarining umumiy soni chekli yoki sanoqli cheksiz bo'lib, ularni 1, 2, 3 kabi natural sonlar bilan sanash mumkin bo'lgan ma'lumotlar."),
    ("Uzluksiz ma'lumotlar: ", "Ma'lum bir oraliqda istalgan qiymatni (jumladan kasr sonlarni) qabul qila oladigan ma'lumotlar (masalan: vaqt, harorat, foiz, vazn).")
]

for bold_p, text_p in concepts:
    add_body(text_p, bold_prefix=bold_p)

add_heading_2("O'lchov shkalalari (Measurement Scales)")
scales_theory = [
    ("1. Nominal shkala: ", "Faqat toifalarni o'z ichiga oladi. Undagi ma'lumotlarni mantiqiy tartiblab yoki saralab bo'lmaydi, ular ustida hech qanday arifmetik amal bajarib bo'lmaydi (masalan: jinsi, fakultet nomi, ID raqamlar)."),
    ("2. Tartib (Ordinal) shkalasi: ", "Obyektlarning nisbiy o'rni va darajasini ko'rsatadi, ammo ular orasidagi masofa aniq miqdoriy ifodalanmaydi (masalan: o'quv kursi [1<2<3<4], musobaqa ligalari [Birinchi < Pro < Superliga], xavf darajasi [Past < O'rta < Yuqori])."),
    ("3. Intervalli (Interval) shkala: ", "Qiymatlar orasidagi farqni hisoblash mumkin bo'lgan shkala, biroq ularning nisbati mutlaq ma'noga ega emas, chunki mutlaq 'nol' (nol deganda xususiyat yo'qligi) mavjud emas. Masalan: gollar farqi (+5, 0, -3), havo harorati (°C)."),
    ("4. Nisbiy (Ratio) shkala: ", "Haqiqiy mutlaq nol boshlang'ich nuqtasi mavjud bo'lgan shkala. Unda qiymatlar farqi ham, ularning o'zaro karrali nisbati ham (masalan, 6 ta gol 3 tadan 2 barobar ko'p) to'liq fizik va matematik ma'noga ega.")
]

for bold_s, text_s in scales_theory:
    add_body(text_s, bold_prefix=bold_s)

doc.add_page_break()

# ==============================================================================
# AMALIY QISM: 1-VAZIFA - FUTBOL STATISTIKASI (TOPSHIRIQ 7)
# ==============================================================================
add_heading_1("II. 1-TANLANGAN TOPSHIRIQ: FUTBOL KLUBLARI STATISTIKASI (7-TOPSHIRIQ)")

add_body("Ushbu bo'limda 'topshiriq_07_futbol_klublar_statistikasi.csv' ma'lumotlar to'plami o'rganilib, topshiriq shartidagi barcha 7 ta amaliy band to'liq bajarildi.")

# 1-band
add_heading_2("1. CSV jadvalining atributlari va ma'lumot turlari")
add_body(f"Jadval {df_football.shape[0]} ta obyekt (qator) va {df_football.shape[1]} ta atributdan (ustun) iborat:")

tbl_attr1 = doc.add_table(rows=1, cols=3)
tbl_attr1.alignment = WD_TABLE_ALIGNMENT.CENTER
set_table_borders(tbl_attr1)

hdr_cells = tbl_attr1.rows[0].cells
hdr_cells[0].text = "№"
hdr_cells[1].text = "Atribut nomi (Ustun)"
hdr_cells[2].text = "Dtype (Dasturdagi turi)"
for cell in hdr_cells:
    set_cell_background(cell, "1E293B")
format_row(tbl_attr1.rows[0], is_header=True)

for idx, (c, dt) in enumerate(df_football.dtypes.items(), 1):
    r_c = tbl_attr1.add_row().cells
    r_c[0].text = str(idx)
    r_c[1].text = str(c)
    r_c[2].text = str(dt)
    bg = "F8FAFC" if idx % 2 == 1 else "FFFFFF"
    for cell in r_c:
        set_cell_background(cell, bg)
    format_row(tbl_attr1.rows[-1])

# 2-band
add_heading_2("2. Tasodifiy 10 ta obyekt (Namuna)")
add_body("Python dasturining `df.sample(10, random_state=42)` funksiyasi yordamida olingan 10 ta tasodifiy obyekt:")

tbl_sample1 = doc.add_table(rows=1, cols=8)
tbl_sample1.alignment = WD_TABLE_ALIGNMENT.CENTER
set_table_borders(tbl_sample1)

sub_cols1 = ["Jamoa_Nomi", "Liga", "Oyinlar", "Galaba", "Durang", "Maglub", "Ochko", "Gollar_Farqi"]
for i, col_name in enumerate(sub_cols1):
    tbl_sample1.rows[0].cells[i].text = col_name
    set_cell_background(tbl_sample1.rows[0].cells[i], "1E293B")
format_row(tbl_sample1.rows[0], is_header=True)

sample10_fb = df_football.sample(10, random_state=42)
for idx, (_, r) in enumerate(sample10_fb.iterrows(), 1):
    row_cells = tbl_sample1.add_row().cells
    row_cells[0].text = str(r["Jamoa_Nomi"])
    row_cells[1].text = str(r["Liga"])
    row_cells[2].text = str(r["Oyinlar_Soni"])
    row_cells[3].text = str(r["Galabalar"])
    row_cells[4].text = str(r["Duranglar"])
    row_cells[5].text = str(r["Maglubiyatlar"])
    row_cells[6].text = str(r["Ochkolar"])
    row_cells[7].text = str(r["Gollar_Farqi"])
    bg = "F8FAFC" if idx % 2 == 1 else "FFFFFF"
    for cell in row_cells:
        set_cell_background(cell, bg)
    format_row(tbl_sample1.rows[-1])

# 3-band
add_heading_2("3. Har bir atribut uchun o'zgaruvchilarni aniqlash")
vars_fb = [
    ("Jamoa_Nomi: ", "Futbol klubining rasmiy nomi bo'lib, jamoalarni bir-biridan ajratuvchi toifaviy o'zgaruvchi."),
    ("Liga: ", "Klub musobaqa olib borayotgan chempionat darajasi (Superliga, Pro-liga, Birinchi liga)."),
    ("Oyinlar_Soni: ", "Chempionat davomida klub o'tkazgan rasmiy uchrashuvlarning butun sonli miqdori."),
    ("Galabalar: ", "Jamoa raqib ustidan g'alaba qozongan o'yinlar miqdori."),
    ("Duranglar: ", "Teng hisobda yakunlangan uchrashuvlar soni."),
    ("Maglubiyatlar: ", "Klub mag'lub bo'lgan uchrashuvlar soni."),
    ("Ochkolar: ", "Jamoaning turnir jadvalidagi o'rnini belgilovchi ochkolar summasi (G'alaba*3 + Durang*1)."),
    ("Urilgan_Gollar: ", "Raqiblar darvozasiga kiritilgan to'plarning jami soni."),
    ("Otkazilgan_Gollar: ", "O'z darvozasidan olib chiqilgan to'plar soni."),
    ("Gollar_Farqi: ", "Urilgan va o'tkazilgan to'plar farqi (+, 0, -)."),
    ("Xavfli_Vaziyatlar: ", "Jamoa hujumlarida yaratilgan darvozaga aniq xavf tug'diruvchi vaziyatlar miqdori."),
    ("Sovrinlar_Oxirgi_5_Yil: ", "Klub so'nggi 5 yil davomida yutgan rasmiy chempionlik va kubok sovrinlari soni.")
]
for bp, tp in vars_fb:
    add_body(tp, bold_prefix=bp)

# 4-band
add_heading_2("4. Har bir atribut uchun umumiy populyatsiyani (bosh to'plam) aniqlash")
pop_fb = [
    ("Jamoa_Nomi: ", "Mamlakatdagi barcha rasmiy ro'yxatdan o'tgan professional futbol klublari."),
    ("Liga: ", "Milliy futbol federatsiyasi tarkibidagi barcha professional va yarim professional divizionlar."),
    ("Oyinlar_Soni: ", "Chempionat nizomi bo'yicha taqvimdagi barcha nazariy o'yinlar miqdori."),
    ("Galabalar, Duranglar, Maglubiyatlar: ", "Barcha o'tkazilgan o'yinlarning mumkin bo'lgan kombinatsiyalari [0; N]."),
    ("Ochkolar: ", "Musobaqada to'planishi mumkin bo'lgan maksimal ochkolar oralig'i [0; 3*N]."),
    ("Urilgan va O'tkazilgan Gollar: ", "Futbol uchrashuvlarida kiritilishi mumkin bo'lgan barcha gollar butun to'plami."),
    ("Gollar_Farqi: ", "Musobaqada yuzaga kelishi mumkin bo'lgan nazariy gollar differensiallari [-M; +M]."),
    ("Xavfli_Vaziyatlar: ", "Chempionat uchrashuvlaridagi barcha xavfli hujumlar va vaziyatlar to'plami."),
    ("Sovrinlar_Oxirgi_5_Yil: ", "So'nggi 5 yillikda professional futbolda o'ynalgan barcha rasmiy kuboklar.")
]
for bp, tp in pop_fb:
    add_body(tp, bold_prefix=bp)

# 5-band
add_heading_2("5. Gipotezalar va ularni ma'lumotlar asosida tekshirish")
hypo_fb_text = [
    ("1-Gipoteza (G'alabalar va Ochkolar): ", "Jamoa qancha ko'p g'alaba qozonsa, to'plagan ochkolari shuncha yuqori bo'ladi. Haqiqiy dataset bo'yicha Pearson korrelyatsiya koeffitsiyenti r = " + f"{df_football['Galabalar'].corr(df_football['Ochkolar']):.4f} ga teng bo'lib, o'zaro deyarli 100% to'g'ri chiziqli qat'iy bog'liqlik borligini isbotladi."),
    ("2-Gipoteza (Xavfli vaziyatlar va Urilgan gollar): ", "Xavfli hujumlar soni ko'p bo'lgan jamoalarda urilgan gollar soni ham sezilarli yuqori bo'ladi. Tekshiruv natijasida korrelyatsiya koeffitsiyenti r = " + f"{df_football['Xavfli_Vaziyatlar'].corr(df_football['Urilgan_Gollar']):.4f} ni tashkil qilib, gipoteza to'liq tasdiqlandi."),
    ("3-Gipoteza (Ligalar bo'yicha ko'rsatkichlar): ", "Ligalar kesimida tahlil qilinganda, Pro liga jamoalari o'rtacha eng ko'p gol (35.6 ta) va ochko (49.4) to'plagan, Superligada raqobat kuchliligi sababli o'rtacha urilgan gol 31.8 tani tashkil etgan.")
]
for bp, tp in hypo_fb_text:
    add_body(tp, bold_prefix=bp)

# 6-band
add_heading_2("6. Ustunlarni diskret va uzluksiz ma'lumotlarga ajratish")
add_body("Ushbu datasetdagi barcha o'lchovlar sanoqli birliklar (butun sonlar va toifalar) bilan ifodalangani uchun quyidagicha ajratiladi:")
add_body("• Diskret ma'lumotlar: Jamoa_Nomi (toifa), Liga (toifa), Oyinlar_Soni, Galabalar, Duranglar, Maglubiyatlar, Ochkolar, Urilgan_Gollar, Otkazilgan_Gollar, Gollar_Farqi, Xavfli_Vaziyatlar, Sovrinlar_Oxirgi_5_Yil (barchasi butun sonli sanoqli qiymatlar).")
add_body("• Uzluksiz ma'lumotlar: Mavjud emas (Agar har bir o'yindagi o'rtacha gollar soni yoki to'p nazorati foizi kiritilsa, ular uzluksiz ustun bo'lar edi).")

# 7-band
add_heading_2("7. Ustunlarni o'lchov shkalalari bo'yicha ajratish")
add_body("• Nominal shkala: Jamoa_Nomi (o'zaro matematik taqqoslab bo'lmaydigan sof nomlar).")
add_body("• Tartib (Ordinal) shkala: Liga (Superliga > Pro-liga > Birinchi liga - ierarxiya mavjud).")
add_body("• Intervalli (Interval) shkala: Gollar_Farqi (farqni hisoblab bo'ladi, ammo '0' gollar yo'qligini emas, balki tenglikni anglatadi; manfiy qiymatlarga ega).")
add_body("• Nisbiy (Ratio) shkala: Oyinlar_Soni, Galabalar, Duranglar, Maglubiyatlar, Ochkolar, Urilgan_Gollar, Otkazilgan_Gollar, Xavfli_Vaziyatlar, Sovrinlar_Oxirgi_5_Yil (mutlaq nol boshlang'ich nuqtasi bor va karrali nisbat ma'noga ega).")

doc.add_page_break()

# ==============================================================================
# AMALIY QISM: 2-VAZIFA - TALABALAR TA'LIM XAVFI (TOPSHIRIQ 3)
# ==============================================================================
add_heading_1("III. 2-TANLANGAN TOPSHIRIQ: TALABALARNING TA'LIM KO'RSATKICHLARI VA XAVF GURUHLARI (3-TOPSHIRIQ)")

add_body("Ushbu bo'limda 'topshiriq_03_talim_talabalar_xavfi.csv' ma'lumotlar to'plami bo'yicha 7 ta amaliy vazifa bajarildi.")

# 1-band
add_heading_2("1. CSV jadvalining atributlari va turlari")
add_body(f"Jadval {df_talaba.shape[0]} nafar talaba (obyekt) va {df_talaba.shape[1]} ta atributdan iborat:")

tbl_attr2 = doc.add_table(rows=1, cols=3)
tbl_attr2.alignment = WD_TABLE_ALIGNMENT.CENTER
set_table_borders(tbl_attr2)

for i, title in enumerate(["№", "Atribut nomi (Ustun)", "Dtype"]):
    tbl_attr2.rows[0].cells[i].text = title
    set_cell_background(tbl_attr2.rows[0].cells[i], "1E293B")
format_row(tbl_attr2.rows[0], is_header=True)

for idx, (c, dt) in enumerate(df_talaba.dtypes.items(), 1):
    r_c = tbl_attr2.add_row().cells
    r_c[0].text = str(idx)
    r_c[1].text = str(c)
    r_c[2].text = str(dt)
    bg = "F8FAFC" if idx % 2 == 1 else "FFFFFF"
    for cell in r_c:
        set_cell_background(cell, bg)
    format_row(tbl_attr2.rows[-1])

# 2-band
add_heading_2("2. Tasodifiy 10 ta obyekt (Namuna)")
add_body("Datasetdan tasodifiy tanlab olingan 10 nafar talaba ko'rsatkichlari:")

tbl_sample2 = doc.add_table(rows=1, cols=7)
tbl_sample2.alignment = WD_TABLE_ALIGNMENT.CENTER
set_table_borders(tbl_sample2)

sub_cols2 = ["Talaba_ID", "Fakultet", "Kurs", "Davomat%", "Oraliq_Baho", "Ishlash_soat", "Xavf_Guruhi"]
for i, col_name in enumerate(sub_cols2):
    tbl_sample2.rows[0].cells[i].text = col_name
    set_cell_background(tbl_sample2.rows[0].cells[i], "1E293B")
format_row(tbl_sample2.rows[0], is_header=True)

sample10_tal = df_talaba.sample(10, random_state=42)
for idx, (_, r) in enumerate(sample10_tal.iterrows(), 1):
    row_cells = tbl_sample2.add_row().cells
    row_cells[0].text = str(r["Talaba_ID"])
    row_cells[1].text = str(r["Fakultet"])
    row_cells[2].text = str(r["Kurs"])
    row_cells[3].text = f"{r['Davomat_Foizi']:.1f}%"
    row_cells[4].text = f"{r['Oraliq_Baho_Ortacha']:.1f}"
    row_cells[5].text = str(r["Haftalik_Ishlash_Soati"])
    row_cells[6].text = str(r["Xavf_Guruhi"])[:15] + "..."
    bg = "F8FAFC" if idx % 2 == 1 else "FFFFFF"
    for cell in row_cells:
        set_cell_background(cell, bg)
    format_row(tbl_sample2.rows[-1])

# 3-band
add_heading_2("3. Har bir atribut uchun o'zgaruvchilarni aniqlash")
vars_tal = [
    ("Talaba_ID: ", "Talabaning HEMIS yoki ta'lim axborot tizimidagi takrorlanmas raqami."),
    ("Fakultet: ", "Talaba ta'lim olayotgan fakultet (AKT, Menejment, Iqtisodiyot va h.k.)."),
    ("Kurs: ", "Talabaning o'quv bosqichi (1, 2, 3, 4-kurs)."),
    ("Davomat_Foizi: ", "Talabaning darslarda qatnashish intizomini ifodalovchi foiz (0-100%)."),
    ("Oraliq_Baho_Ortacha: ", "Oraliq nazoratlardan to'plangan o'rtacha o'zlashtirish balli."),
    ("Mustaqil_Ish_Foizi: ", "Mustaqil ta'lim topshiriqlarining o'z vaqtida bajarilish ko'rsatkichi (%)."),
    ("Ijtimoiy_Yordam: ", "Talabaga moddiy rag'bat yoki ijtimoiy imtiyoz berilganligi ('Ha' / 'Yo'q')."),
    ("Yotoqxonada_Turadi: ", "Talabaning talabalar turar joyida istiqomat qilishi ('Ha' / 'Yo'q')."),
    ("Haftalik_Ishlash_Soati: ", "Talabaning o'qishdan tashqari haftalik mehnat bilan bandlik vaqti (soat)."),
    ("Xavf_Guruhi: ", "Talabaning akademik qarzdor bo'lib qolish xavfi darajasi (Past, O'rta, Yuqori).")
]
for bp, tp in vars_tal:
    add_body(tp, bold_prefix=bp)

# 4-band
add_heading_2("4. Har bir atribut uchun umumiy populyatsiyani (bosh to'plam) aniqlash")
pop_tal = [
    ("Talaba_ID: ", "Oliy ta'lim muassasasida ro'yxatga olingan barcha talabalar kontingenti."),
    ("Fakultet: ", "Universitetdagi barcha mavjud fakultetlar va yo'nalishlar majmuasi."),
    ("Kurs: ", "Bakalavriat ta'lim bosqichidagi barcha kurslar [1-kursdan 4-kursgacha]."),
    ("Davomat_Foizi: ", "Barcha talabalarning semestr davomidagi nazariy davomat ko'rsatkichlari [0.0% - 100.0%]."),
    ("Oraliq_Baho_Ortacha: ", "Barcha talabalarning oraliq baholash natijalari [0.0 - 100.0 ball]."),
    ("Mustaqil_Ish_Foizi: ", "O'quv rejasidagi barcha mustaqil ish topshiriqlarining qamrovi [0% - 100%]."),
    ("Ijtimoiy_Yordam / Yotoqxona: ", "Universitetdagi ijtimoiy himoyaga muhtoj va turar joy bilan ta'minlangan barcha talabalar toifasi."),
    ("Haftalik_Ishlash_Soati: ", "Ishlaydigan va ishlamaydigan barcha talabalarning haftalik darsdan tashqari bandlik soatlari majmuasi."),
    ("Xavf_Guruhi: ", "Barcha talabalarning o'zlashtirish va o'qishdan qolish xavfi bo'yicha to'liq toifalari.")
]
for bp, tp in pop_tal:
    add_body(tp, bold_prefix=bp)

# 5-band
add_heading_2("5. Gipotezalar va ularni ma'lumotlar asosida tekshirish")
hypo_tal_text = [
    ("1-Gipoteza (Davomat va Mustaqil ish foizi): ", "Darsga intizom bilan muntazam qatnashuvchi talabalar mustaqil ta'lim topshiriqlarini ham yuqori ko'rsatkichda topshiradi. Haqiqiy dataset bo'yicha hisoblangan Pearson korrelyatsiya koeffitsiyenti r = " + f"{df_talaba['Davomat_Foizi'].corr(df_talaba['Mustaqil_Ish_Foizi']):.4f} ga teng bo'lib, o'zaro kuchli to'g'ridan-to'g'ri chiziqli bog'liqlik borligini isbotladi."),
    ("2-Gipoteza (Haftalik ishlash soati va Xavf guruhi): ", "Haftasiga ko'proq soat darsdan tashqari ishlaydigan talabalarda akademik xavf darajasi sezilarli ortadi. Haqiqiy ma'lumotlar tahlili: 'Past xavf' (xavfsiz) talabalar haftasiga o'rtacha 4.7 soat, 'O'rta xavf' talabalar 14.4 soat, 'Yuqori xavf' guruhidagi talabalar esa haftasiga o'rtacha 20.5 soat mehnat qiladi. Gipoteza to'liq tasdiqlandi!"),
    ("3-Gipoteza (Davomat ko'rsatkichi va Xavf guruhi): ", "Davomat foizi yuqori bo'lgan (o'rtacha 91.6%) talabalar xavfsiz (past xavf) guruhini tashkil etadi, o'rtacha davomati 71.0% gacha pasaygan talabalar esa to'g'ridan-to'g'ri 'Yuqori xavf' guruhiga tushadi.")
]
for bp, tp in hypo_tal_text:
    add_body(tp, bold_prefix=bp)

# 6-band
add_heading_2("6. Ustunlarni diskret va uzluksiz ma'lumotlarga ajratish")
add_body("• Diskret ma'lumotlar: Talaba_ID (toifa), Fakultet (toifa), Kurs (butun son: 1, 2, 3, 4), Ijtimoiy_Yordam (Ha/Yo'q), Yotoqxonada_Turadi (Ha/Yo'q), Haftalik_Ishlash_Soati (butun soatlar), Xavf_Guruhi (toifa).")
add_body("• Uzluksiz ma'lumotlar: Davomat_Foizi (0% dan 100% gacha bo'lgan uzluksiz o'nli kasr foiz), Oraliq_Baho_Ortacha (haqiqiy uzluksiz ball), Mustaqil_Ish_Foizi (uzluksiz foiz ko'rsatkichi).")

# 7-band
add_heading_2("7. Ustunlarni o'lchov shkalalari bo'yicha ajratish")
add_body("• Nominal shkala: Talaba_ID (identifikator), Fakultet (toifalar), Ijtimoiy_Yordam (Ha/Yo'q), Yotoqxonada_Turadi (Ha/Yo'q).")
add_body("• Tartib (Ordinal) shkala: Kurs (1 < 2 < 3 < 4 kurs ierarxiyasi), Xavf_Guruhi (Past xavf < O'rta xavf < Yuqori xavf tartibi).")
add_body("• Intervalli shkala: Sof intervalli shkala mavjud emas (chunki barcha raqamli ko'rsatkichlar tabiiy nolga ega).")
add_body("• Nisbiy (Ratio) shkala: Davomat_Foizi (0% mutlaq nol), Oraliq_Baho_Ortacha (0 ball mutlaq nol), Mustaqil_Ish_Foizi (0%), Haftalik_Ishlash_Soati (0 soat).")

doc.add_page_break()

# ==============================================================================
# DASTURIY TA'MINOT VA INTERFEYS
# ==============================================================================
add_heading_1("IV. DASTURIY TA'MINOT: PYTHON VA FOYDALANUVCHI INTERFEYSI")

add_body("Ushbu amaliy mashg'ulot talabiga asosan («Barcha topshiriqlar Python dasturlash tilida, bitta dasturda hamda qulay interfeysga ega ko‘rinishda tuzilsin»), '2_amaliy_dastur.py' nomli to'liq yagona dastur ishlab chiqildi. Dastur quyidagi xususiyatlarga ega:")

feats = [
    ("1. Ikki tomonlama interfeys: ", "Dastur zamonaviy grafik oynali GUI (Tkinter) interfeysiga hamda konsol (CLI) menyusiga ega."),
    ("2. Interaktiv dataset almashtirish: ", "Foydalanuvchi bir marta bosish orqali 7-topshiriq (Futbol) yoki 3-topshiriq (Talabalar xavfi) datasetlarini tanlashi mumkin."),
    ("3. Alohida va umumiy hisobot: ", "Har bir 1-7 vazifani alohida ko'rish yoki '★ TO'LIQ HISOBOTNI CHIQARISH' tugmasi orqali barchasini birdaniga konsolga yoki matn oynasiga chiqarish imkoniyati mavjud."),
    ("4. Gipotezalarni matematik tekshirish: ", "Dastur gipotezalarni shunchaki matn sifatida ko'rsatmaydi, balki pandas kutubxonasi yordamida korrelyatsiya (Pearson r) va guruhlar bo'yicha o'rtacha qiymatlarni hisoblab, ilmiy isbotlab beradi."),
    ("5. Nusxa olish (Copy to clipboard): ", "Chiqarilgan har qanday natijani buferga nusxalash tugmasi integratsiya qilingan.")
]
for bp, tp in feats:
    add_body(tp, bold_prefix=bp)

add_heading_2("Dasturni ishga tushirish bo'yicha yo'riqnoma")
add_body("• Grafik interfeysda (GUI) ochish uchun:")
add_body("   python 2_amaliy_dastur.py", italic=True)
add_body("• Konsol (CLI) menyusida ochish uchun:")
add_body("   python 2_amaliy_dastur.py --cli", italic=True)

doc.add_page_break()

# ==============================================================================
# XULOSA VA ADABIYOTLAR
# ==============================================================================
add_heading_1("V. XULOSA")

add_body("2-amaliy mashg'ulotni bajarish davomida quyidagi muhim ilmiy-amaliy natijalar olindi:")
add_body("1. Ma'lumotlar to'plami strukturasini tahlil qilishda obyekt, atribut, o'zgaruvchi, bosh to'plam va tanlanma tushunchalari amaliy ma'lumotlar misolida chuqur o'rganildi.")
add_body("2. Ma'lumotlarni to'g'ri qayta ishlash va mashinali o'qitish modellariga tayyorlash uchun ularning diskret yoki uzluksiz ekanligini, shuningdek o'lchov shkalasini (nominal, tartib, intervalli, nisbiy) to'g'ri aniqlash o'ta muhim ahamiyatga ega ekanligi isbotlandi.")
add_body("3. 1-amaliyotda yaratilgan 2 ta dataset (Futbol klublari va Talabalar ta'lim xavfi) bo'yicha tuzilgan barcha gipotezalar Python (pandas) vositalari bilan to'liq tekshirildi va ijobiy korrelyatsiyalar aniqlandi.")
add_body("4. Topshiriq shartiga muvofiq, bitta yagona interaktiv dastur yaratilib, foydalanuvchiga qulay grafik va konsol interfeysi taqdim etildi.")

add_heading_1("VI. FOYDALANILGAN ADABIYOTLAR")
sources = [
    "1. Han, J., Kamber, M., & Pei, J. (2011). Data Mining: Concepts and Techniques. Morgan Kaufmann.",
    "2. Tan, P. N., Steinbach, M., & Kumar, V. (2018). Introduction to Data Mining. Pearson Education.",
    "3. McKinney, W. (2022). Python for Data Analysis: Data Wrangling with Pandas, NumPy, and Jupyter (3rd ed.). O'Reilly Media.",
    "4. Stevens, S. S. (1946). On the theory of scales of measurement. Science, 103(2684), 677-680.",
    "5. 2-amaliy mashg'ulot uslubiy qo'llanmasi: Ma'lumotlar to'plami va uning atributlari. Metama'lumotlar."
]
for src in sources:
    add_body(src)

# Faylni saqlash
docx_path = os.path.join(CURRENT_DIR, "2-amaliy_mashg'ulot_hisoboti.docx")
doc.save(docx_path)
print(f"Hujjat saqlandi: {docx_path}")

# ==============================================================================
# JUPYTER NOTEBOOK YARATISH (2-amaliy_mashgulot.ipynb)
# ==============================================================================

nb_cells = [
    {
        "cell_type": "markdown",
        "metadata": {},
        "source": [
            "# 2-Amaliy Mashg'ulot: Ma'lumotlar to'plami va uning atributlari. Metama'lumotlar\n",
            "**Fan:** Data Mining  \n",
            "**Mavzu:** CSV fayllaridagi ma'lumotlar to'plamlarini tahlil qilish, atributlar, umumiy populyatsiya, gipotezalar, diskret/uzluksiz turlari va o'lchov shkalalari.  \n",
            "**Tanlangan topshiriqlar:**  \n",
            "1. 7-topshiriq: Futbol klublari statistikasi (`topshiriq_07_futbol_klublar_statistikasi.csv`)  \n",
            "2. 3-topshiriq: Talabalar ta'lim monitoringi va chetlatilish xavfi (`topshiriq_03_talim_talabalar_xavfi.csv`)"
        ]
    },
    {
        "cell_type": "code",
        "execution_count": None,
        "metadata": {},
        "outputs": [],
        "source": [
            "# Kerakli kutubxonalarni import qilish\n",
            "import os\n",
            "import pandas as pd\n",
            "import numpy as np\n",
            "import matplotlib.pyplot as plt\n",
            "\n",
            "# Datasetlar joylashuvi\n",
            "DATASETS_DIR = 'datasets'\n",
            "print('Kutubxonalar muvaffaqiyatli yuklandi.')"
        ]
    },
    {
        "cell_type": "markdown",
        "metadata": {},
        "source": [
            "## 1-BO'LIM: 7-topshiriq - Futbol klublari statistikasi\n",
            "### 1. CSV jadvalining atributlarini chop etish"
        ]
    },
    {
        "cell_type": "code",
        "execution_count": None,
        "metadata": {},
        "outputs": [],
        "source": [
            "# 1. CSV faylini o'qish va atributlarni chiqarish\n",
            "df_football = pd.read_csv(os.path.join(DATASETS_DIR, 'topshiriq_07_futbol_klublar_statistikasi.csv'))\n",
            "print(f'Jadval hajmi: {df_football.shape[0]} qator x {df_football.shape[1]} ustun\\n')\n",
            "print('Atributlar ro\\'yxati va turlari:')\n",
            "display(df_football.dtypes.to_frame('Dtype'))"
        ]
    },
    {
        "cell_type": "markdown",
        "metadata": {},
        "source": [
            "### 2. Tasodifiy 10 ta obyektni chop etish"
        ]
    },
    {
        "cell_type": "code",
        "execution_count": None,
        "metadata": {},
        "outputs": [],
        "source": [
            "# 2. Tasodifiy 10 ta obyekt (Namuna)\n",
            "sample_football = df_football.sample(n=10, random_state=42)\n",
            "display(sample_football)"
        ]
    },
    {
        "cell_type": "markdown",
        "metadata": {},
        "source": [
            "### 3 & 4. O'zgaruvchilar va Umumiy Populyatsiya (Bosh to'plam) tavsifi\n",
            "| Atribut nomi | O'zgaruvchi tavsifi | Umumiy populyatsiya (Bosh to'plam) |\n",
            "| :--- | :--- | :--- |\n",
            "| `Jamoa_Nomi` | Klubning rasmiy nomi (unikal toifaviy belgi) | Mamlakatdagi barcha professional futbol klublari |\n",
            "| `Liga` | Klub o'ynaydigan liga (Superliga, Pro, Birinchi) | Milliy futbol assotsiatsiyasidagi barcha ligalar |\n",
            "| `Oyinlar_Soni` | Mavsumda o'tkazilgan jami uchrashuvlar | Chempionat nizomidagi barcha o'yinlar miqdori |\n",
            "| `Galabalar` | Jamoa g'alaba qozongan o'yinlar soni | Barcha o'yinlardagi potensial g'alabalar [0, N] |\n",
            "| `Duranglar` | Durang qayd etilgan uchrashuvlar soni | Barcha o'yinlardagi potensial duranglar |\n",
            "| `Maglubiyatlar` | Mag'lubiyatlar soni | Barcha o'yinlardagi potensial mag'lubiyatlar |\n",
            "| `Ochkolar` | To'plangan jami ochkolar (3*G + 1*D) | Nazariy to'planishi mumkin bo'lgan ochkolar [0, 3*N] |\n",
            "| `Urilgan_Gollar` | Raqib darvozasiga urilgan to'plar | Barcha o'yinlarda kiritilishi mumkin bo'lgan gollar |\n",
            "| `Otkazilgan_Gollar` | O'z darvozasidan o'tkazilgan gollar | Barcha o'tkazib yuborilishi mumkin bo'lgan gollar |\n",
            "| `Gollar_Farqi` | Urilgan va o'tkazilgan gollar farqi | Musobaqadagi barcha nazariy differensiallar [-M, +M] |\n",
            "| `Xavfli_Vaziyatlar` | Yaratilgan xavfli golli vaziyatlar soni | Chempionatdagi barcha xavfli hujumlar to'plami |\n",
            "| `Sovrinlar_Oxirgi_5_Yil` | So'nggi 5 yilda yutilgan kuboklar | So'nggi 5 yillikdagi barcha musobaqa sovrinlari |"
        ]
    },
    {
        "cell_type": "markdown",
        "metadata": {},
        "source": [
            "### 5. Gipotezalar va ularni ma'lumotlar asosida tekshirish"
        ]
    },
    {
        "cell_type": "code",
        "execution_count": None,
        "metadata": {},
        "outputs": [],
        "source": [
            "# Gipoteza 1: G'alabalar va Ochkolar o'rtasidagi bog'liqlik\n",
            "r1 = df_football['Galabalar'].corr(df_football['Ochkolar'])\n",
            "print(f'H1 tekshiruvi: G\\'alabalar va Ochkolar korrelyatsiyasi = {r1:.4f}')\n",
            "\n",
            "# Gipoteza 2: Xavfli vaziyatlar va Urilgan gollar bog'liqligi\n",
            "r2 = df_football['Xavfli_Vaziyatlar'].corr(df_football['Urilgan_Gollar'])\n",
            "print(f'H2 tekshiruvi: Xavfli vaziyatlar va Urilgan gollar korrelyatsiyasi = {r2:.4f}')\n",
            "\n",
            "# Gipoteza 3: Ligalar bo'yicha o'rtacha ko'rsatkichlar\n",
            "print('\\nH3 tekshiruvi: Ligalar kesimidagi o\\'rtacha ko\\'rsatkichlar:')\n",
            "display(df_football.groupby('Liga')[['Urilgan_Gollar', 'Otkazilgan_Gollar', 'Ochkolar']].mean().round(1))"
        ]
    },
    {
        "cell_type": "markdown",
        "metadata": {},
        "source": [
            "### 6 & 7. Diskret/Uzluksiz turlar hamda O'lchov shkalalari\n",
            "* **Diskret ma'lumotlar:** Jamoa_Nomi, Liga, Oyinlar_Soni, Galabalar, Duranglar, Maglubiyatlar, Ochkolar, Urilgan_Gollar, Otkazilgan_Gollar, Gollar_Farqi, Xavfli_Vaziyatlar, Sovrinlar_Oxirgi_5_Yil.  \n",
            "* **Uzluksiz ma'lumotlar:** Mavjud emas (barcha qiymatlar butun sonli va toifaviy).  \n",
            "\n",
            "* **Nominal shkala:** `Jamoa_Nomi` (Nomlar, toifalar).  \n",
            "* **Tartib (Ordinal) shkala:** `Liga` (Superliga > Pro-liga > Birinchi liga).  \n",
            "* **Intervalli (Interval) shkala:** `Gollar_Farqi` (Farq hisoblanadi, 0 qiymat tenglikni bildiradi, manfiy qiymatlar mavjud).  \n",
            "* **Nisbiy (Ratio) shkala:** `Oyinlar_Soni`, `Galabalar`, `Duranglar`, `Maglubiyatlar`, `Ochkolar`, `Urilgan_Gollar`, `Otkazilgan_Gollar`, `Xavfli_Vaziyatlar`, `Sovrinlar_Oxirgi_5_Yil`."
        ]
    },
    {
        "cell_type": "markdown",
        "metadata": {},
        "source": [
            "---\n",
            "## 2-BO'LIM: 3-topshiriq - Talabalar ta'lim monitoringi va chetlatilish xavfi\n",
            "### 1. CSV jadvalining atributlarini chop etish"
        ]
    },
    {
        "cell_type": "code",
        "execution_count": None,
        "metadata": {},
        "outputs": [],
        "source": [
            "# 1. Talabalar datasetini o'qish va atributlarni chiqarish\n",
            "df_talaba = pd.read_csv(os.path.join(DATASETS_DIR, 'topshiriq_03_talim_talabalar_xavfi.csv'))\n",
            "print(f'Jadval hajmi: {df_talaba.shape[0]} qator x {df_talaba.shape[1]} ustun\\n')\n",
            "display(df_talaba.dtypes.to_frame('Dtype'))"
        ]
    },
    {
        "cell_type": "markdown",
        "metadata": {},
        "source": [
            "### 2. Tasodifiy 10 ta obyektni chop etish"
        ]
    },
    {
        "cell_type": "code",
        "execution_count": None,
        "metadata": {},
        "outputs": [],
        "source": [
            "# 2. Tasodifiy 10 ta talaba namunasi\n",
            "display(df_talaba.sample(n=10, random_state=42))"
        ]
    },
    {
        "cell_type": "markdown",
        "metadata": {},
        "source": [
            "### 3 & 4. O'zgaruvchilar va Umumiy Populyatsiya (Bosh to'plam) tavsifi\n",
            "| Atribut nomi | O'zgaruvchi tavsifi | Umumiy populyatsiya (Bosh to'plam) |\n",
            "| :--- | :--- | :--- |\n",
            "| `Talaba_ID` | Talabaning tizimdagi unikal identifikatori | Universitetdagi barcha ro'yxatdan o'tgan talabalar |\n",
            "| `Fakultet` | Tahsil olayotgan fakultet nomi | OTMdagi barcha mavjud fakultetlar va yo'nalishlar |\n",
            "| `Kurs` | Talabaning o'quv kursi (1, 2, 3, 4) | Bakalavriatning barcha o'quv kurslari [1-4] |\n",
            "| `Davomat_Foizi` | Darslarga qatnashish foizi (0-100%) | Semestr davomidagi nazariy davomat oralig'i [0-100%] |\n",
            "| `Oraliq_Baho_Ortacha` | Oraliq nazorat imtihonlari o'rtacha balli | Baholash tizimidagi barcha oraliq ballar [0-100] |\n",
            "| `Mustaqil_Ish_Foizi` | Mustaqil ishlarning bajarilish foizi | Barcha belgilangan mustaqil ish topshiriqlari |\n",
            "| `Ijtimoiy_Yordam` | Ijtimoiy imtiyoz mavjudligi (Ha/Yo'q) | OTMdagi ijtimoiy himoyaga muhtoj barcha talabalar |\n",
            "| `Yotoqxonada_Turadi` | TTJda istiqomat qilishi (Ha/Yo'q) | Yotoqxona bilan ta'minlangan barcha talabalar |\n",
            "| `Haftalik_Ishlash_Soati` | Darsdan tashqari mehnat soatlari | Ishlaydigan va ishlamaydigan talabalarning bandlik soatlari |\n",
            "| `Xavf_Guruhi` | Akademik xavf toifasi (Past, O'rta, Yuqori) | Barcha talabalarning xavf guruhlari bo'yicha to'plami |"
        ]
    },
    {
        "cell_type": "markdown",
        "metadata": {},
        "source": [
            "### 5. Gipotezalar va ularni ma'lumotlar asosida tekshirish"
        ]
    },
    {
        "cell_type": "code",
        "execution_count": None,
        "metadata": {},
        "outputs": [],
        "source": [
            "# Gipoteza 1: Davomat va Mustaqil ish foizi o'rtasidagi bog'liqlik\n",
            "r_dav = df_talaba['Davomat_Foizi'].corr(df_talaba['Mustaqil_Ish_Foizi'])\n",
            "print(f'H1 tekshiruvi: Davomat va Mustaqil ish foizi korrelyatsiyasi = {r_dav:.4f}')\n",
            "\n",
            "# Gipoteza 2: Haftalik ishlash soati va Xavf guruhi\n",
            "print('\\nH2 tekshiruvi: Xavf guruhlari bo\\'yicha o\\'rtacha ishlash soati:')\n",
            "display(df_talaba.groupby('Xavf_Guruhi')['Haftalik_Ishlash_Soati'].agg(['mean', 'count']).round(1))\n",
            "\n",
            "# Gipoteza 3: Davomat ko'rsatkichi va Xavf guruhi\n",
            "print('\\nH3 tekshiruvi: Xavf guruhlari bo\\'yicha o\\'rtacha davomat foizi:')\n",
            "display(df_talaba.groupby('Xavf_Guruhi')['Davomat_Foizi'].agg(['mean', 'count']).round(1))"
        ]
    },
    {
        "cell_type": "markdown",
        "metadata": {},
        "source": [
            "### 6 & 7. Diskret/Uzluksiz turlar hamda O'lchov shkalalari\n",
            "* **Diskret ma'lumotlar:** `Talaba_ID`, `Fakultet`, `Kurs`, `Ijtimoiy_Yordam`, `Yotoqxonada_Turadi`, `Haftalik_Ishlash_Soati`, `Xavf_Guruhi`.  \n",
            "* **Uzluksiz ma'lumotlar:** `Davomat_Foizi`, `Oraliq_Baho_Ortacha`, `Mustaqil_Ish_Foizi`.  \n",
            "\n",
            "* **Nominal shkala:** `Talaba_ID`, `Fakultet`, `Ijtimoiy_Yordam`, `Yotoqxonada_Turadi`.  \n",
            "* **Tartib (Ordinal) shkala:** `Kurs` (1 < 2 < 3 < 4), `Xavf_Guruhi` (Past xavf < O'rta xavf < Yuqori xavf).  \n",
            "* **Intervalli shkala:** Sof intervalli ustun mavjud emas.  \n",
            "* **Nisbiy (Ratio) shkala:** `Davomat_Foizi`, `Oraliq_Baho_Ortacha`, `Mustaqil_Ish_Foizi`, `Haftalik_Ishlash_Soati`."
        ]
    }
]

nb_dict = {
    "cells": nb_cells,
    "metadata": {
        "kernelspec": {
            "display_name": "Python 3",
            "language": "python",
            "name": "python3"
        },
        "language_info": {
            "name": "python",
            "version": "3.12"
        }
    },
    "nbformat": 4,
    "nbformat_minor": 5
}

nb_path = os.path.join(CURRENT_DIR, "2-amaliy_mashgulot.ipynb")
with open(nb_path, "w", encoding="utf-8") as f:
    json.dump(nb_dict, f, ensure_ascii=False, indent=2)

print(f"Jupyter Notebook saqlandi: {nb_path}")
